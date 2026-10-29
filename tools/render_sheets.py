"""Render the weekly report with LibreOffice: one PDF page and one PNG per sheet; optionally check totals.

Usage: python tools/render_sheets.py web/latest/Cedarline_Weekly_Report_*.xlsx --out web/latest/preview [--check-recalc]
Needs soffice (libreoffice-calc-nogui), pdfinfo and pdftoppm (poppler-utils); CI installs them, the owner's
machine does not have them.
"""

from __future__ import annotations

import argparse
import math
import os
import shutil
import subprocess
import sys
import tempfile
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from buildreport.config import Config, load_config
from buildreport.isoweek import IsoWeek
from buildreport.metrics import compute
from buildreport.workbook import data, formula_cells, runinfo

ROOT = Path(__file__).resolve().parents[1]
PAGES = ("summary", "branches", "products", "data", "run-info")
MAX_PNG_BYTES = 400_000
PDF_FILTER = 'pdf:calc_pdf_Export:{"SinglePageSheets":{"type":"boolean","value":"true"}}'
RECALC_FILTER = "xlsx:Calc MS Excel 2007 XML"
PROFILE = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


def fail(message: str) -> int:
    print(f"render_sheets: {message}", file=sys.stderr)
    return 1


def make_render_copy(source: Path, target: Path, rows: int = 40) -> int:
    """A copy whose Data table holds the header and the first ``rows`` rows plus a note row."""
    book = load_workbook(source)
    sheet = book[data.SHEET]
    total = sheet.max_row - 1
    if total > rows:
        sheet.delete_rows(rows + 2, sheet.max_row - rows - 1)
        sheet.cell(
            row=rows + 2, column=1, value=f"first {rows} of {total:,} rows; the full table is in the workbook"
        )
        table = sheet.tables[data.TABLE]
        table.ref = f"A1:H{rows + 1}"
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref
    book.save(target)
    return total


def write_profile(profile: Path) -> str:
    """A LibreOffice user profile that always recalculates formulas when it loads xlsx and ods files."""
    user = profile / "user"
    user.mkdir(parents=True, exist_ok=True)
    (user / "registrymodifications.xcu").write_text(PROFILE, encoding="utf-8")
    return profile.resolve().as_uri()


def _soffice(soffice: str, profile_uri: str) -> list[str]:
    return [soffice, "--headless", "--norestore", f"-env:UserInstallation={profile_uri}"]


def pdf_command(soffice: str, profile_uri: str, workbook: Path, outdir: Path) -> list[str]:
    return [
        *_soffice(soffice, profile_uri),
        "--convert-to",
        PDF_FILTER,
        "--outdir",
        str(outdir),
        str(workbook),
    ]


def recalc_command(soffice: str, profile_uri: str, workbook: Path, outdir: Path) -> list[str]:
    return [
        *_soffice(soffice, profile_uri),
        "--convert-to",
        RECALC_FILTER,
        "--outdir",
        str(outdir),
        str(workbook),
    ]


def _run(command: list[str]) -> str:
    """Run a command with a 10 minute limit; its stdout, or a RuntimeError that carries its stderr."""
    completed = subprocess.run(command, capture_output=True, text=True, timeout=600)
    if completed.returncode != 0:
        raise RuntimeError(f"{command[0]} failed ({completed.returncode}): {completed.stderr.strip()}")
    return completed.stdout


def pdf_pages(pdf: Path) -> int:
    for line in _run(["pdfinfo", str(pdf)]).splitlines():
        if line.startswith("Pages:"):
            return int(line.split(":", 1)[1])
    raise RuntimeError(f"pdfinfo printed no page count for {pdf}")


def rename_pages(folder: Path) -> list[Path]:
    """pdftoppm writes sheet-1.png (or sheet-01.png); give the pages the sheet names in order."""
    pages = sorted(folder.glob("sheet-*.png"), key=lambda path: int(path.stem.rsplit("-", 1)[1]))
    renamed = []
    for path, name in zip(pages, PAGES, strict=False):
        target = folder / f"{name}.png"
        path.replace(target)
        renamed.append(target)
    return renamed


def oversized(pngs: Iterable[Path], limit: int = MAX_PNG_BYTES) -> list[str]:
    """The names of the images larger than ``limit`` bytes."""
    return [path.name for path in pngs if path.stat().st_size > limit]


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return float(value)


def check_recalc(original: Path, recalculated: Path, cfg: Config) -> list[str]:
    """Every generated formula, recalculated by LibreOffice, against the engine (relative 1e-9)."""
    book = load_workbook(original)
    period = IsoWeek.parse(str(book[runinfo.SHEET]["B3"].value))
    result = compute(data.read_records(book), cfg, period)
    values = load_workbook(recalculated, data_only=True)
    problems = []
    for cell in formula_cells(cfg, result):
        raw = values[cell.sheet][cell.ref].value
        actual = _number(raw)
        if cell.expected is None:
            ok = raw in (None, "")
        else:
            ok = actual is not None and math.isclose(actual, cell.expected, rel_tol=1e-9, abs_tol=1e-12)
        if not ok:
            problems.append(
                f"{cell.sheet}!{cell.ref}: {cell.formula} gives {raw!r}, the engine says {cell.expected!r}"
            )
    return problems


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("workbooks", nargs="+", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--check-recalc", action="store_true")
    parser.add_argument("--config", type=Path, default=ROOT / "report.yaml")
    parser.add_argument("--rows", type=int, default=40)
    args = parser.parse_args(argv)
    existing = [path for path in args.workbooks if path.is_file()]
    if not existing:
        return fail("no workbook to render (the pattern matched no file)")
    if len(existing) > 1:
        return fail(f"expected one workbook, got {len(existing)}")
    workbook = existing[0]
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None or shutil.which("pdftoppm") is None or shutil.which("pdfinfo") is None:
        return fail(
            "LibreOffice (soffice) and poppler-utils (pdftoppm, pdfinfo) are needed; the render runs in CI"
        )
    try:
        profile_dir = Path(tempfile.mkdtemp(prefix="lo-profile-", dir=os.environ.get("RUNNER_TEMP") or None))
    except OSError as exc:
        return fail(f"cannot create the LibreOffice profile folder: {exc}")
    try:
        profile_uri = write_profile(profile_dir)
        with tempfile.TemporaryDirectory(prefix="render-") as folder:
            work = Path(folder)
            render_copy = work / "render.xlsx"
            total = make_render_copy(workbook, render_copy, args.rows)
            _run(pdf_command(soffice, profile_uri, render_copy, work))
            pdf = work / "render.pdf"
            pages = pdf_pages(pdf)
            if pages != len(PAGES):
                return fail(f"the preview PDF has {pages} pages, expected {len(PAGES)} (one per sheet)")
            _run(
                [
                    "pdftoppm",
                    "-png",
                    "-scale-to-x",
                    "1440",
                    "-scale-to-y",
                    "-1",
                    str(pdf),
                    str(work / "sheet"),
                ]
            )
            pngs = rename_pages(work)
            if len(pngs) != len(PAGES):
                return fail(f"pdftoppm wrote {len(pngs)} images, expected {len(PAGES)}")
            too_big = oversized(pngs)
            if too_big:
                return fail(f"over {MAX_PNG_BYTES:,} bytes: {', '.join(too_big)}")
            args.out.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(pdf, args.out / "report-preview.pdf")
            for path in pngs:
                shutil.copyfile(path, args.out / path.name)
        if args.check_recalc:
            recalc_dir = args.out / "recalc"
            recalc_dir.mkdir(parents=True, exist_ok=True)
            _run(recalc_command(soffice, profile_uri, workbook, recalc_dir))
            problems = check_recalc(workbook, recalc_dir / workbook.name, load_config(args.config))
            for problem in problems:
                print(problem, file=sys.stderr)
            if problems:
                return fail(f"{len(problems)} recalculated totals differ from the engine")
            print("recalculated totals match the engine")
    except (RuntimeError, subprocess.SubprocessError, OSError) as exc:
        return fail(str(exc))
    finally:
        shutil.rmtree(profile_dir, ignore_errors=True)
    print(
        f"rendered {workbook.name}: {len(PAGES)} pages, data table cut to {min(args.rows, total)} of {total:,} rows"
    )
    for name in PAGES:
        path = args.out / f"{name}.png"
        print(f"  {path.name:<14} {path.stat().st_size:>9,} bytes")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
