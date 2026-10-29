import hashlib
import json
import shutil
import subprocess
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import urlparse
from urllib.request import url2pathname

import pytest
from openpyxl import load_workbook
from PIL import Image, ImageDraw

from conftest import STEM

DROP = object()


@pytest.fixture()
def report_dir(tmp_path, demo_run):
    folder = tmp_path / "report"
    shutil.copytree(demo_run["dir"], folder)
    return folder


@pytest.fixture()
def old_target(tmp_path):
    target = tmp_path / "latest"
    (target / "preview").mkdir(parents=True)
    (target / ".gitkeep").write_bytes(b"")
    (target / "Cedarline_Weekly_Report_2025-W37.xlsx").write_bytes(b"old")
    (target / "preview" / "summary.png").write_bytes(b"old")
    return target


def snapshot(folder):
    """Every path below the folder with its bytes (None for a folder)."""
    return {
        path.relative_to(folder).as_posix(): path.read_bytes() if path.is_file() else None
        for path in sorted(folder.rglob("*"))
    }


def edit_manifest(report_dir, content):
    """Rewrite the run's manifest: raw bytes, or a dict of changes (DROP removes the key)."""
    path = report_dir / f"{STEM}.manifest.json"
    if isinstance(content, bytes):
        path.write_bytes(content)
        return
    manifest = json.loads(path.read_text(encoding="utf-8"))
    for key, value in content.items():
        if value is DROP:
            del manifest[key]
        else:
            manifest[key] = value
    path.write_text(json.dumps(manifest), encoding="utf-8")


def test_publish_latest_replaces_the_folder(tool, report_dir, old_target, tmp_path, capsys):
    target = old_target
    assert tool("publish_latest").main([str(report_dir), str(target)]) == 0
    assert capsys.readouterr().out == "2025-W38\n"
    assert sorted(path.name for path in target.iterdir()) == sorted(
        [".gitkeep", f"{STEM}.xlsx", f"{STEM}.eml", f"{STEM}.manifest.json", "manifest.json"]
    )
    assert (target / "manifest.json").read_bytes() == (report_dir / f"{STEM}.manifest.json").read_bytes()
    assert (target / f"{STEM}.xlsx").read_bytes() == (report_dir / f"{STEM}.xlsx").read_bytes()
    assert sorted(path.name for path in tmp_path.iterdir()) == ["latest", "report"]


def test_publish_latest_refuses_a_blocked_or_missing_report(tool, report_dir, tmp_path):
    publish = tool("publish_latest").main
    assert publish([str(tmp_path / "empty"), str(tmp_path / "latest")]) == 1
    manifest_path = report_dir / f"{STEM}.manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["status"] = "blocked"
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    assert publish([str(report_dir), str(tmp_path / "latest")]) == 1
    manifest["status"] = "built"
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    (report_dir / f"{STEM}.eml").unlink()
    assert publish([str(report_dir), str(tmp_path / "latest")]) == 1


def test_render_copy_cuts_the_data_table(tool, demo_run, tmp_path):
    render = tool("render_sheets")
    target = tmp_path / "render.xlsx"
    assert render.make_render_copy(demo_run["xlsx"], target, rows=40) == 2945
    book = load_workbook(target)
    data = book["Data"]
    assert data.max_row == 42
    assert data["A42"].value == "first 40 of 2,945 rows; the full table is in the workbook"
    table = data.tables["tblWeekly"]
    assert table.ref == "A1:H41" and table.autoFilter.ref == "A1:H41"
    assert book.sheetnames == ["Summary", "Branches", "Products", "Data", "Run info"]
    assert len(book["Summary"]._charts) == 3
    assert book["Branches"]["D16"].value == "=SUM(D4:D15)"


def test_profile_sets_recalculation_to_always(tool, tmp_path):
    uri = tool("render_sheets").write_profile(tmp_path / "lo-profile")
    assert uri.startswith("file:")
    text = (tmp_path / "lo-profile" / "user" / "registrymodifications.xcu").read_text(encoding="utf-8")
    for name in ("OOXMLRecalcMode", "ODFRecalcMode"):
        assert f'<prop oor:name="{name}" oor:op="fuse"><value>0</value></prop>' in text


def test_soffice_command_lines(tool, tmp_path):
    render = tool("render_sheets")
    pdf = render.pdf_command("soffice", "file:///tmp/lo-profile", tmp_path / "render.xlsx", tmp_path)
    assert pdf[:4] == ["soffice", "--headless", "--norestore", "-env:UserInstallation=file:///tmp/lo-profile"]
    assert pdf[4:6] == [
        "--convert-to",
        'pdf:calc_pdf_Export:{"SinglePageSheets":{"type":"boolean","value":"true"}}',
    ]
    recalc = render.recalc_command(
        "soffice", "file:///tmp/lo-profile", tmp_path / "report.xlsx", tmp_path / "recalc"
    )
    assert recalc[4:6] == ["--convert-to", "xlsx:Calc MS Excel 2007 XML"]


def test_rename_pages(tool, tmp_path):
    for number in range(1, 6):
        (tmp_path / f"sheet-{number}.png").write_bytes(b"png")
    names = [path.name for path in tool("render_sheets").rename_pages(tmp_path)]
    assert names == ["summary.png", "branches.png", "products.png", "data.png", "run-info.png"]


def test_check_recalc_compares_every_formula_cell(tool, cfg, demo_run, demo_result, tmp_path):
    from buildreport.workbook import formula_cells

    render = tool("render_sheets")
    simulated = load_workbook(demo_run["xlsx"])
    cells = formula_cells(cfg, demo_result)
    for cell in cells:
        simulated[cell.sheet][cell.ref] = "" if cell.expected is None else cell.expected
    good = tmp_path / "good.xlsx"
    simulated.save(good)
    assert render.check_recalc(demo_run["xlsx"], good, cfg) == []
    simulated["Branches"]["D16"] = cells[0].expected * 1.01
    bad = tmp_path / "bad.xlsx"
    simulated.save(bad)
    problems = render.check_recalc(demo_run["xlsx"], bad, cfg)
    assert len(problems) == 1 and problems[0].startswith("Branches!D16: =SUM(D4:D15) gives")


def test_render_needs_libreoffice(tool, demo_run, tmp_path, monkeypatch, capsys):
    render = tool("render_sheets")
    monkeypatch.setattr(render.shutil, "which", lambda name: None)
    assert render.main([str(demo_run["xlsx"]), "--out", str(tmp_path / "preview")]) == 1
    assert "LibreOffice" in capsys.readouterr().err


def test_render_needs_exactly_one_existing_workbook(tool, tmp_path, capsys):
    render = tool("render_sheets")
    assert (
        render.main([str(tmp_path / "Cedarline_Weekly_Report_*.xlsx"), "--out", str(tmp_path / "preview")])
        == 1
    )
    assert "no workbook" in capsys.readouterr().err


@pytest.mark.parametrize(("damage", "word"), [("flip", "sha256"), ("append", "bytes")])
def test_publish_latest_verifies_the_manifest_digests(tool, report_dir, old_target, capsys, damage, word):
    xlsx = report_dir / f"{STEM}.xlsx"
    blob = bytearray(xlsx.read_bytes())
    if damage == "flip":
        blob[len(blob) // 2] ^= 0x01
    else:
        blob += b"x"
    xlsx.write_bytes(bytes(blob))
    before = snapshot(old_target)
    assert tool("publish_latest").main([str(report_dir), str(old_target)]) == 1
    err = capsys.readouterr().err
    assert err.startswith("publish_latest: ") and f"{STEM}.xlsx" in err and word in err
    assert snapshot(old_target) == before


@pytest.mark.parametrize(
    "content",
    [
        b"{not json",
        b"\xff\xfe\x00",
        b"[1, 2]",
        {"period": DROP},
        {"status": DROP},
        {"outputs": DROP},
        {"period": None},
        {"outputs": "nope"},
        {"outputs": [1]},
        {"outputs": [{"bytes": 1, "sha256": "x"}]},
        {"outputs": [{"file": "a.xlsx", "bytes": "9", "sha256": "x"}]},
        {"outputs": [{"file": "a.xlsx", "bytes": 9, "sha256": None}]},
    ],
)
def test_publish_latest_refuses_a_malformed_manifest(tool, report_dir, old_target, capsys, content):
    edit_manifest(report_dir, content)
    before = snapshot(old_target)
    assert tool("publish_latest").main([str(report_dir), str(old_target)]) == 1
    assert capsys.readouterr().err.startswith("publish_latest: ")
    assert snapshot(old_target) == before


def test_publish_latest_refuses_two_manifests(tool, report_dir, old_target, capsys):
    shutil.copyfile(report_dir / f"{STEM}.manifest.json", report_dir / "Other.manifest.json")
    before = snapshot(old_target)
    assert tool("publish_latest").main([str(report_dir), str(old_target)]) == 1
    assert "found 2" in capsys.readouterr().err
    assert snapshot(old_target) == before


@pytest.mark.parametrize(
    "name", ["../evil.xlsx", "..", ".", "", "sub/report.xlsx", "sub\\report.xlsx", "/abs.xlsx"]
)
def test_publish_latest_refuses_names_that_leave_the_folder(
    tool, report_dir, old_target, tmp_path, capsys, name
):
    evil = b"evil"
    (tmp_path / "evil.xlsx").write_bytes(evil)
    (report_dir / "sub").mkdir()
    (report_dir / "sub" / "report.xlsx").write_bytes(evil)
    entry = {"file": name, "bytes": len(evil), "sha256": hashlib.sha256(evil).hexdigest()}
    edit_manifest(report_dir, {"outputs": [entry]})
    before = snapshot(old_target)
    assert tool("publish_latest").main([str(report_dir), str(old_target)]) == 1
    assert "plain file name" in capsys.readouterr().err
    assert snapshot(old_target) == before


def test_publish_latest_refuses_overlapping_folders(tool, report_dir, tmp_path, capsys):
    publish = tool("publish_latest").main
    before = snapshot(report_dir)
    siblings = sorted(path.name for path in tmp_path.iterdir())
    assert publish([str(report_dir), str(report_dir)]) == 1
    assert publish([str(report_dir), str(report_dir / ".." / "report")]) == 1
    assert publish([str(report_dir), str(tmp_path)]) == 1
    assert publish([str(report_dir), str(report_dir / "latest")]) == 1
    assert capsys.readouterr().err.count("same folder or one holds the other") == 4
    assert snapshot(report_dir) == before
    assert sorted(path.name for path in tmp_path.iterdir()) == siblings


def test_publish_latest_unlinks_symlinks_and_refuses_a_symlinked_target(
    tool, report_dir, old_target, tmp_path
):
    publish = tool("publish_latest").main
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "keep.txt").write_bytes(b"keep")
    try:
        (old_target / "linked").symlink_to(outside, target_is_directory=True)
        (tmp_path / "alias").symlink_to(old_target, target_is_directory=True)
    except (OSError, NotImplementedError):
        pytest.skip("symlinks are not available here")
    before = snapshot(old_target)
    assert publish([str(report_dir), str(tmp_path / "alias")]) == 1
    assert snapshot(old_target) == before
    assert publish([str(report_dir), str(old_target)]) == 0
    assert not (old_target / "linked").is_symlink() and not (old_target / "linked").exists()
    assert (outside / "keep.txt").read_bytes() == b"keep"


def test_publish_latest_stages_the_copy(tool, report_dir, old_target, tmp_path, monkeypatch, capsys):
    before = snapshot(old_target)
    real, calls = Path.write_bytes, []

    def flaky(self, data):
        calls.append(self.name)
        if len(calls) == 2:
            raise OSError("disk full")
        return real(self, data)

    monkeypatch.setattr(Path, "write_bytes", flaky)
    assert tool("publish_latest").main([str(report_dir), str(old_target)]) == 1
    assert "disk full" in capsys.readouterr().err
    assert snapshot(old_target) == before
    assert sorted(path.name for path in tmp_path.iterdir()) == ["latest", "report"]


def _sheet_like_image(width: int = 1440, height: int = 2000) -> Image.Image:
    """An RGB page: white, a coloured header band, a grid of grey rules and a gradient block of many colours."""
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle([0, 0, width, 80], fill=(31, 54, 92))
    for y in range(120, height, 24):
        draw.line([(40, y), (width - 40, y)], fill=(210, 210, 210))
    for x in range(0, 400):
        draw.line([(40 + x, 200), (40 + x, 600)], fill=(x % 256, (x * 3) % 256, 255 - x % 256))
    return image


def test_shrink_stores_a_palette_png_under_the_limit(tool, tmp_path):
    render = tool("render_sheets")
    png = tmp_path / "data.png"
    _sheet_like_image().save(png, format="PNG")
    before = png.stat().st_size
    after = render.shrink(png)
    assert after == png.stat().st_size < before
    assert after <= render.MAX_PNG_BYTES
    with Image.open(png) as image:
        assert image.mode == "P" and image.size == (1440, 2000)
        assert image.convert("RGB").getpixel((5, 5)) == (31, 54, 92)  # the header colour survives
        assert image.convert("RGB").getpixel((1400, 1900)) == (255, 255, 255)


def test_shrink_tries_smaller_palettes_only_while_over_the_limit(tool, tmp_path, monkeypatch):
    render = tool("render_sheets")
    png = tmp_path / "run-info.png"
    _sheet_like_image(600, 400).save(png, format="PNG")
    sizes = []
    original = render.Image.Image.quantize

    def counting(self, *args, **kwargs):
        sizes.append(kwargs["colors"])
        return original(self, *args, **kwargs)

    monkeypatch.setattr(render.Image.Image, "quantize", counting)
    render.shrink(png)
    assert sizes == [256]
    sizes.clear()
    render.shrink(png, limit=1)
    assert sizes == [256, 64, 16]


def test_oversized_pages_are_named(tool, tmp_path):
    render = tool("render_sheets")
    small, big = tmp_path / "summary.png", tmp_path / "data.png"
    small.write_bytes(b"x" * render.MAX_PNG_BYTES)
    big.write_bytes(b"x" * (render.MAX_PNG_BYTES + 1))
    assert render.oversized([small, big]) == ["data.png"]
    assert render.oversized([small, big], limit=10) == ["summary.png", "data.png"]


def test_pdf_pages_has_a_timeout_and_reports_stderr(tool, tmp_path, monkeypatch):
    render = tool("render_sheets")
    calls, outcome = [], {}

    def run(command, **kwargs):
        calls.append((command, kwargs))
        return subprocess.CompletedProcess(command, outcome["code"], outcome["out"], outcome["err"])

    monkeypatch.setattr(render.subprocess, "run", run)
    outcome.update(code=0, out="Title: x\nPages:          5\n", err="")
    assert render.pdf_pages(tmp_path / "render.pdf") == 5
    assert calls[0][0][0] == "pdfinfo" and calls[0][1]["timeout"] == 600
    outcome.update(code=0, out="Title: x\n", err="")
    with pytest.raises(RuntimeError, match="no page count"):
        render.pdf_pages(tmp_path / "render.pdf")
    outcome.update(code=1, out="", err="Syntax Error: bad xref table\n")
    with pytest.raises(RuntimeError, match=r"pdfinfo failed \(1\): Syntax Error: bad xref table"):
        render.pdf_pages(tmp_path / "render.pdf")


@pytest.fixture()
def fake_render(tool, tmp_path, monkeypatch):
    """render_sheets.main with the external programs replaced; LibreOffice and poppler never run."""
    render = tool("render_sheets")
    runner = tmp_path / "runner"
    runner.mkdir()
    state = SimpleNamespace(render=render, runner=runner, profiles=[], png_size=lambda number: 1000)
    monkeypatch.setenv("RUNNER_TEMP", str(runner))
    monkeypatch.setattr(render.shutil, "which", lambda name: name)

    def run(command):
        if command[0] == "soffice":
            uri = next(item for item in command if item.startswith("-env:UserInstallation=")).split("=", 1)[1]
            profile = Path(url2pathname(urlparse(uri).path))
            assert profile.is_dir() and profile.parent == runner and profile.name.startswith("lo-profile-")
            state.profiles.append(profile)
            (Path(command[command.index("--outdir") + 1]) / "render.pdf").write_bytes(b"%PDF")
            return ""
        if command[0] == "pdfinfo":
            return "Pages:          5\n"
        assert command[0] == "pdftoppm"
        prefix = Path(command[-1])
        for number in range(1, 6):
            (prefix.parent / f"{prefix.name}-{number}.png").write_bytes(b"x" * state.png_size(number))
        return ""

    monkeypatch.setattr(render, "_run", run)
    monkeypatch.setattr(render, "shrink", lambda png, limit=render.MAX_PNG_BYTES: png.stat().st_size)
    return state


def test_render_uses_a_fresh_profile_per_run_and_removes_it(fake_render, demo_run, tmp_path):
    main, workbook = fake_render.render.main, str(demo_run["xlsx"])
    assert main([workbook, "--out", str(tmp_path / "first")]) == 0
    assert main([workbook, "--out", str(tmp_path / "second")]) == 0
    assert len(set(fake_render.profiles)) == 2
    assert not any(path.exists() for path in fake_render.profiles)
    assert list(fake_render.runner.iterdir()) == []
    assert sorted(path.name for path in (tmp_path / "first").iterdir()) == [
        "branches.png",
        "data.png",
        "products.png",
        "report-preview.pdf",
        "run-info.png",
        "summary.png",
    ]


def test_render_checks_the_image_size_before_copying_anything(fake_render, demo_run, tmp_path, capsys):
    fake_render.png_size = lambda number: fake_render.render.MAX_PNG_BYTES + 1 if number == 4 else 1000
    out = tmp_path / "preview"
    assert fake_render.render.main([str(demo_run["xlsx"]), "--out", str(out)]) == 1
    assert "over 400,000 bytes: data.png" in capsys.readouterr().err
    assert not out.exists()
    assert list(fake_render.runner.iterdir()) == []
