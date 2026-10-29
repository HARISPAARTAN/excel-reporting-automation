import io
import math
import random
from datetime import datetime

import pytest
from openpyxl import load_workbook

from buildreport.fixtures import (
    FORMATS,
    content_fingerprint,
    generate,
    growth,
    make_spec,
    mix_for,
    normal,
    seasonality,
    write_fixtures,
)
from buildreport.isoweek import IsoWeek

W38 = IsoWeek(2025, 38)


@pytest.fixture(scope="module")
def demo_files(cfg):
    return {item.name: item.data for item in generate(cfg, W38, 8, deviations="demo")}


@pytest.mark.parametrize(
    ("week", "factor"),
    [
        (IsoWeek(2025, 1), 0.86),
        (IsoWeek(2025, 5), 0.90),
        (IsoWeek(2025, 6), 0.96),
        (IsoWeek(2025, 22), 1.00),
        (IsoWeek(2025, 23), 1.02),
        (IsoWeek(2025, 34), 0.95),
        (IsoWeek(2025, 38), 1.00),
        (IsoWeek(2025, 47), 1.03),
        (IsoWeek(2025, 48), 1.08),
        (IsoWeek(2025, 51), 1.24),
        (IsoWeek(2025, 52), 1.05),
        (IsoWeek(2026, 53), 0.92),
    ],
)
def test_seasonality_table(week, factor):
    assert seasonality(week) == factor


def test_growth_starts_at_the_base_monday():
    assert growth(IsoWeek(2025, 1)) == 1.0
    assert growth(W38) == pytest.approx(1.03 ** (259 / 365.25))


def test_mix_sums_to_one_and_follows_the_season():
    for fmt in FORMATS:
        for week in (IsoWeek(2025, 5), IsoWeek(2025, 25), W38):
            assert math.fsum(mix_for(fmt, week).values()) == pytest.approx(1.0)
    summer, autumn = mix_for("Neighbourhood", IsoWeek(2025, 25)), mix_for("Neighbourhood", W38)
    assert summer["Hot Drinks"] < autumn["Hot Drinks"]
    assert summer["Cold Drinks and Grocery"] > autumn["Cold Drinks and Grocery"]


def test_normal_draws_are_truncated():
    rng = random.Random("cedarline:test")
    draws = [normal(rng, 0.05, 0.85, 1.15) for _ in range(2000)]
    assert all(0.85 <= value <= 1.15 for value in draws)
    assert math.fsum(draws) / len(draws) == pytest.approx(1.0, abs=0.01)
    assert all(0.94 <= normal(rng, 0.02) <= 1.06 for _ in range(500))


def test_design_targets_for_2025_w38(cfg):
    specs = [make_spec(cfg, code, W38) for code in cfg.branch_codes]
    sales = [math.fsum(line.net_sales_eur for line in spec.sales) for spec in specs]
    assert 270_000 <= math.fsum(sales) <= 330_000
    waste = [
        math.fsum(line.waste_eur for line in spec.sales) / total
        for spec, total in zip(specs, sales, strict=True)
    ]
    assert all(0.030 <= value <= 0.085 for value in waste)
    assert any(value > 0.06 for value in waste)
    labour = [
        math.fsum(line.cost_eur for line in spec.labour) / total
        for spec, total in zip(specs, sales, strict=True)
    ]
    assert all(0.22 <= value <= 0.38 for value in labour)


def test_generation_is_deterministic_and_seeded(cfg):
    first = generate(cfg, W38, 1, deviations="none")
    second = generate(cfg, W38, 1, deviations="none")
    assert [item.data for item in first] == [item.data for item in second]
    other = generate(cfg, W38, 1, deviations="none", seed=2208)
    assert content_fingerprint(first[0].data) != content_fingerprint(other[0].data)


def test_demo_profile_file_names(demo_files):
    names = list(demo_files)
    assert len(names) == 96 and names == sorted(names)
    assert "CL-11_2025-W28.xlsx" in names and "CL-11_2025-W38.xlsx" not in names
    assert "CL-10_2025-W37.xlsx" in names and "CL-10_2025-W37_resend.xlsx" in names
    assert "CL-12_2025-W36.xlsx" not in names


def test_demo_deviations_inside_the_files(demo_files):
    def book(name):
        return load_workbook(io.BytesIO(demo_files[name]))

    assert book("CL-03_2025-W38.xlsx").sheetnames == ["Cover", "Sales wk38", "Labour"]
    header = [cell.value for cell in book("CL-05_2025-W38.xlsx")["Sales"][1]]
    assert header == ["Product group", "Net sales EUR", "Units sold", "Notes", "Waste units", "Waste EUR"]
    blank = book("CL-08_2025-W38.xlsx")["Cover"]
    assert blank["A5"].value == "ISO week" and blank["B5"].value is None
    assert book("CL-11_2025-W28.xlsx")["Cover"]["B5"].value == "2025-W38"
    first, resend = book("CL-10_2025-W37.xlsx"), book("CL-10_2025-W37_resend.xlsx")
    assert first["Cover"]["B7"].value == datetime(2025, 9, 15, 8, 40)
    assert resend["Cover"]["B7"].value == datetime(2025, 9, 15, 14, 12)
    assert resend["Sales"]["A2"].value == "Bread"
    assert resend["Sales"]["B2"].value == round(first["Sales"]["B2"].value * 1.035, 2)
    assert resend["Sales"]["B3"].value == first["Sales"]["B3"].value


def test_deviations_outside_the_generated_weeks_are_skipped(cfg):
    names = [item.name for item in generate(cfg, W38, 1, deviations="demo")]
    assert len(names) == 12
    assert not any(name.endswith("_resend.xlsx") for name in names)


def test_the_misnamed_file_avoids_a_collision(cfg):
    names = {item.name for item in generate(cfg, W38, 11, deviations="demo")}
    assert {"CL-11_2025-W28.xlsx", "CL-11_2025-W28_late.xlsx"} <= names


def test_write_refuses_to_overwrite_without_force(cfg, tmp_path):
    files = generate(cfg, W38, 1, deviations="none")
    write_fixtures(files, tmp_path)
    with pytest.raises(FileExistsError, match="use --force"):
        write_fixtures(files, tmp_path)
    write_fixtures(files, tmp_path, force=True)
    assert len(list(tmp_path.glob("*.xlsx"))) == 12


def test_committed_demo_set_matches_the_generator(tool):
    assert tool("check_fixtures").main([]) == 0


def test_files_are_small(demo_files):
    assert all(3_000 < len(data) < 20_000 for data in demo_files.values())
