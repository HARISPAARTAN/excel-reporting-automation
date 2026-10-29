import email
import email.policy
import json
import smtplib
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from conftest import ROOT, STEM

CONFIG = ROOT / "report.yaml"


def load_manifest(path):
    return json.loads(path.read_text(encoding="utf-8"))


def test_version(run_cli):
    result = run_cli(["--version"])
    assert result.exit_code == 0 and result.output == "build-report 0.1.0\n"


def test_help_lists_the_commands(run_cli):
    result = run_cli(["--help"])
    assert result.exit_code == 0
    for command in ("run", "validate", "generate-fixtures"):
        assert command in result.output


def test_validate_text_without_a_period(run_cli, demo_dir):
    result = run_cli(["validate", "--input", demo_dir, "--config", CONFIG])
    assert result.exit_code == 0, result.output
    lines = result.stdout.splitlines()
    assert len(lines) == 5 and all(line.startswith("warning") for line in lines)
    assert not any("MISSING_BRANCH" in line for line in lines)
    assert "checked 96 files: 0 errors, 5 warnings" in result.stderr


def test_validate_json_with_a_period(run_cli, demo_dir):
    result = run_cli(
        ["validate", "--input", demo_dir, "--period", "2025-W38", "--config", CONFIG, "--format", "json"]
    )
    assert result.exit_code == 0, result.output
    findings = json.loads(result.stdout)
    assert sorted(item["code"] for item in findings) == [
        "DUPLICATE_REPLACED",
        "EXTRA_COLUMN",
        "MISSING_BRANCH",
        "SHEET_RENAMED",
        "WEEK_MISSING",
        "WEEK_NAME_MISMATCH",
    ]
    assert list(findings[0]) == ["severity", "code", "file", "message"]


def test_validate_period_last_uses_the_epoch(run_cli, demo_dir):
    result = run_cli(
        ["validate", "--input", demo_dir, "--period", "last", "--config", CONFIG, "--format", "json"]
    )
    assert result.exit_code == 0
    assert any(
        "2025-W36" in item["message"]
        for item in json.loads(result.stdout)
        if item["code"] == "MISSING_BRANCH"
    )


def test_validate_strict_and_errors(run_cli, demo_dir, scenario):
    assert run_cli(["validate", "--input", demo_dir, "--config", CONFIG, "--strict"]).exit_code == 3
    (scenario / "CL-09_2025-W38.xlsx").write_bytes(b"not a zip file")
    result = run_cli(["validate", "--input", scenario, "--config", CONFIG])
    assert result.exit_code == 3 and "FILE_UNREADABLE" in result.stdout


@pytest.mark.parametrize("period", ["2025-W60", "W38", "2025-38", "next"])
def test_a_malformed_period_is_a_usage_error(run_cli, demo_dir, period, tmp_path):
    result = run_cli(["run", "--period", period, "--input", demo_dir, "--out", tmp_path, "--config", CONFIG])
    assert result.exit_code == 2
    assert "is not an ISO week such as 2025-W38, or last" in result.output
    assert list(tmp_path.iterdir()) == []


def test_run_dry_run_writes_the_three_outputs(demo_run):
    assert demo_run["xlsx"].is_file() and demo_run["eml"].is_file() and demo_run["manifest"].is_file()
    manifest = load_manifest(demo_run["manifest"])
    assert (manifest["status"], manifest["exit_code"], manifest["notify"]) == ("built", 0, "dry-run")
    assert len(manifest["findings"]) == 6 and {item["severity"] for item in manifest["findings"]} == {
        "warning"
    }
    assert [item["file"] for item in manifest["outputs"]] == [f"{STEM}.xlsx", f"{STEM}.eml"]
    assert "6 warnings" in demo_run["result"].stderr
    assert sorted(path.name for path in demo_run["dir"].iterdir()) == sorted(
        [f"{STEM}.xlsx", f"{STEM}.eml", f"{STEM}.manifest.json"]
    )


def test_manifest_hashes_match_the_files(demo_run):
    import hashlib

    manifest = load_manifest(demo_run["manifest"])
    for item in manifest["outputs"]:
        data = (demo_run["dir"] / item["file"]).read_bytes()
        assert item["bytes"] == len(data) and item["sha256"] == hashlib.sha256(data).hexdigest()


def test_eml_carries_the_workbook(demo_run):
    message = email.message_from_bytes(demo_run["eml"].read_bytes(), policy=email.policy.default)
    attachment = next(message.iter_attachments())
    assert attachment.get_content() == demo_run["xlsx"].read_bytes()
    assert message["X-Build-Report-Mode"] == "dry-run"
    assert message.get_body(preferencelist=("plain",)) is not None
    assert message.get_body(preferencelist=("html",)) is not None


def test_kpis_in_the_workbook_equal_the_manifest(demo_run):
    manifest = load_manifest(demo_run["manifest"])
    ws = load_workbook(demo_run["xlsx"])["Summary"]
    for row, (key, entry) in enumerate(manifest["kpis"].items(), start=5):
        assert ws[f"B{row}"].value == entry["label"], key
        assert ws[f"C{row}"].value == pytest.approx(entry["value"], rel=1e-12)
        assert ws[f"D{row}"].value == pytest.approx(entry["previous"], rel=1e-12)


def test_strict_run_blocks_and_removes_stale_outputs(run_cli, demo_dir, tmp_path):
    (tmp_path / f"{STEM}.xlsx").write_bytes(b"stale")
    (tmp_path / f"{STEM}.eml").write_bytes(b"stale")
    args = ["run", "--period", "2025-W38", "--input", demo_dir, "--out", tmp_path, "--config", CONFIG]
    result = run_cli([*args, "--strict", "--notify", "dry-run"])
    assert result.exit_code == 3
    assert sorted(path.name for path in tmp_path.iterdir()) == [f"{STEM}.manifest.json"]
    manifest = load_manifest(tmp_path / f"{STEM}.manifest.json")
    assert (manifest["status"], manifest["exit_code"], manifest["outputs"]) == ("blocked", 3, [])


def test_notify_none_writes_no_eml(run_cli, demo_dir, tmp_path):
    (tmp_path / f"{STEM}.eml").write_bytes(b"stale")
    result = run_cli(
        ["run", "--period", "2025-W38", "--input", demo_dir, "--out", tmp_path, "--config", CONFIG]
    )
    assert result.exit_code == 0
    assert sorted(path.name for path in tmp_path.iterdir()) == sorted(
        [f"{STEM}.xlsx", f"{STEM}.manifest.json"]
    )
    assert load_manifest(tmp_path / f"{STEM}.manifest.json")["notify"] == "none"


class FakeSMTP:
    instances: list["FakeSMTP"] = []
    fail: Exception | None = None

    def __init__(self, host, port, timeout=None, context=None):
        self.host, self.port, self.timeout = host, port, timeout
        self.calls = []
        type(self).instances.append(self)

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def starttls(self, context=None):
        self.calls.append("starttls")

    def login(self, user, password):
        self.calls.append("login")

    def send_message(self, message, from_addr=None, to_addrs=None):
        if type(self).fail is not None:
            raise type(self).fail
        self.calls.append(("send", from_addr, tuple(to_addrs)))


@pytest.fixture()
def fake_smtp(monkeypatch):
    class Fake(FakeSMTP):
        instances = []
        fail = None

    monkeypatch.setattr(smtplib, "SMTP", Fake)
    monkeypatch.setattr(smtplib, "SMTP_SSL", Fake)
    return Fake


SMTP_ENV = {
    "BUILD_REPORT_SMTP_HOST": "smtp.cedarline.example",
    "BUILD_REPORT_SMTP_USER": "reporting",
    "BUILD_REPORT_SMTP_PASSWORD": "secret-from-env",
}


def test_smtp_sends_to_to_and_cc(run_cli, demo_dir, tmp_path, fake_smtp):
    args = [
        "run",
        "--period",
        "2025-W38",
        "--input",
        demo_dir,
        "--out",
        tmp_path,
        "--config",
        CONFIG,
        "--notify",
        "smtp",
    ]
    result = run_cli(args, env=SMTP_ENV)
    assert result.exit_code == 0, result.output
    (client,) = fake_smtp.instances
    assert (client.host, client.port, client.timeout) == ("smtp.cedarline.example", 587, 30)
    assert client.calls == [
        "starttls",
        "login",
        (
            "send",
            "reporting@cedarline.example",
            ("head-office@cedarline.example", "operations@cedarline.example", "finance@cedarline.example"),
        ),
    ]
    message = email.message_from_bytes((tmp_path / f"{STEM}.eml").read_bytes(), policy=email.policy.default)
    assert message["X-Build-Report-Mode"] == "smtp"


def test_smtp_failure_keeps_the_report_and_exits_4(run_cli, demo_dir, tmp_path, fake_smtp):
    fake_smtp.fail = smtplib.SMTPException("connection refused by the test")
    args = [
        "run",
        "--period",
        "2025-W38",
        "--input",
        demo_dir,
        "--out",
        tmp_path,
        "--config",
        CONFIG,
        "--notify",
        "smtp",
    ]
    result = run_cli(args, env=SMTP_ENV)
    assert result.exit_code == 4
    assert (tmp_path / f"{STEM}.xlsx").is_file() and (tmp_path / f"{STEM}.eml").is_file()
    assert load_manifest(tmp_path / f"{STEM}.manifest.json")["exit_code"] == 4


def test_smtp_without_host_is_a_usage_error_before_any_output(run_cli, demo_dir, tmp_path):
    out = tmp_path / "out"
    args = [
        "run",
        "--period",
        "2025-W38",
        "--input",
        demo_dir,
        "--out",
        out,
        "--config",
        CONFIG,
        "--notify",
        "smtp",
    ]
    result = run_cli(args)
    assert result.exit_code == 2
    assert "BUILD_REPORT_SMTP_HOST" in result.stderr
    assert not out.exists()


def test_generate_fixtures_refuses_to_overwrite(run_cli, tmp_path):
    args = [
        "generate-fixtures",
        "--weeks",
        "1",
        "--end-week",
        "2025-W38",
        "--out",
        tmp_path,
        "--config",
        CONFIG,
    ]
    first = run_cli(args)
    assert first.exit_code == 0, first.output
    assert len(list(tmp_path.glob("*.xlsx"))) == 12
    second = run_cli(args)
    assert second.exit_code == 2 and "--force" in second.stderr
    assert run_cli([*args, "--force"]).exit_code == 0


def test_generate_fixtures_last_and_demo(run_cli, tmp_path):
    result = run_cli(
        ["generate-fixtures", "--weeks", "8", "--deviations", "demo", "--out", tmp_path, "--config", CONFIG]
    )
    assert result.exit_code == 0
    names = sorted(path.name for path in tmp_path.glob("*.xlsx"))
    assert len(names) == 96 and "CL-11_2025-W28.xlsx" in names and "CL-10_2025-W37_resend.xlsx" in names


def test_invalid_configuration_is_exit_2(run_cli, demo_dir, tmp_path):
    bad = tmp_path / "report.yaml"
    bad.write_text(CONFIG.read_text(encoding="utf-8") + "colour: blue\n", encoding="utf-8")
    result = run_cli(["validate", "--input", demo_dir, "--config", bad])
    assert result.exit_code == 2 and "colour" in result.stderr


def test_invalid_source_date_epoch_is_exit_2(run_cli, demo_dir, tmp_path):
    args = ["run", "--period", "2025-W38", "--input", demo_dir, "--out", tmp_path, "--config", CONFIG]
    result = run_cli(args, env={"SOURCE_DATE_EPOCH": "yesterday"})
    assert result.exit_code == 2 and "SOURCE_DATE_EPOCH" in result.stderr


def test_missing_input_folder_is_exit_2(run_cli, tmp_path):
    result = run_cli(["validate", "--input", tmp_path / "nowhere", "--config", CONFIG])
    assert result.exit_code == 2 and "does not exist" in result.stderr


def test_log_file_is_appended(run_cli, demo_dir, tmp_path):
    log = tmp_path / "logs" / "build-report.log"
    for _ in range(2):
        assert (
            run_cli(["--log-file", log, "validate", "--input", demo_dir, "--config", CONFIG]).exit_code == 0
        )
    assert log.read_text(encoding="utf-8").count("checked 96 files") == 2


def test_an_unexpected_error_is_exit_1(run_cli, demo_dir, tmp_path, monkeypatch):
    import buildreport.cli

    def boom(ctx):
        raise RuntimeError("boom")

    monkeypatch.setattr(buildreport.cli, "render_bytes", boom)
    args = ["run", "--period", "2025-W38", "--input", demo_dir, "--out", tmp_path, "--config", CONFIG]
    result = run_cli(args)
    assert result.exit_code == 1 and "unexpected error: RuntimeError: boom" in result.stderr


def test_python_dash_m():
    completed = subprocess.run(
        [sys.executable, "-m", "buildreport", "--version"], capture_output=True, text=True, check=True
    )
    assert completed.stdout == "build-report 0.1.0\n"
