import io
import os
import subprocess
import sys
import zipfile

from openpyxl import load_workbook

from buildreport.fixtures import content_fingerprint

from conftest import EPOCH, ROOT, STEM


def build(out, folder="data/incoming"):
    """One report run in a fresh Python process (own hash seed), as the README documents it."""
    env = {
        key: value for key, value in os.environ.items() if not key.startswith(("GITHUB_", "BUILD_REPORT_"))
    }
    env["SOURCE_DATE_EPOCH"] = EPOCH
    command = [
        sys.executable,
        "-m",
        "buildreport",
        "run",
        "--period",
        "2025-W38",
        "--input",
        str(folder),
        "--out",
        str(out),
    ]
    completed = subprocess.run(
        [*command, "--notify", "dry-run"], cwd=ROOT, env=env, capture_output=True, text=True
    )
    assert completed.returncode == 0, completed.stderr
    return {suffix: (out / f"{STEM}{suffix}").read_bytes() for suffix in (".xlsx", ".manifest.json", ".eml")}


def test_two_runs_give_byte_identical_outputs(tmp_path):
    first = build(tmp_path / "a")
    second = build(tmp_path / "b")
    for suffix in (".xlsx", ".manifest.json", ".eml"):
        assert first[suffix] == second[suffix], suffix


def test_every_zip_member_carries_the_epoch(demo_run):
    data = demo_run["xlsx"].read_bytes()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        infos = archive.infolist()
        assert infos and all(info.date_time == (2025, 9, 22, 6, 0, 0) for info in infos)
        assert all(info.create_system == 0 and info.external_attr == 0o600 << 16 for info in infos)
        assert all(info.compress_type == zipfile.ZIP_DEFLATED for info in infos)
        core = archive.read("docProps/core.xml").decode("utf-8")
    assert core.count("2025-09-22T06:00:00Z") == 2


def test_manifest_and_eml_times_follow_the_epoch(demo_run):
    assert '"generated_at": "2025-09-22T06:00:00Z"' in demo_run["manifest"].read_text(encoding="utf-8")
    assert b"Date: Mon, 22 Sep 2025 06:00:00 +0000\r\n" in demo_run["eml"].read_bytes()


def test_one_changed_input_cell_changes_the_workbook(run_cli, scenario, tmp_path):
    def run(out):
        args = [
            "run",
            "--period",
            "2025-W38",
            "--input",
            scenario,
            "--out",
            out,
            "--config",
            ROOT / "report.yaml",
        ]
        assert run_cli(args).exit_code == 0
        return (out / f"{STEM}.xlsx").read_bytes()

    def units(data):
        rows = load_workbook(io.BytesIO(data))["Data"].iter_rows(min_row=2, values_only=True)
        return next(
            row[6]
            for row in rows
            if row[0] == "2025-W38" and row[2] == "CL-01" and row[4] == "Bread" and row[5] == "units_sold"
        )

    before = run(tmp_path / "before")
    path = scenario / "CL-01_2025-W38.xlsx"
    book = load_workbook(path)
    book["Sales"]["C2"] = book["Sales"]["C2"].value + 1
    book.save(path)
    after = run(tmp_path / "after")
    assert after != before
    assert units(after) == units(before) + 1


def test_generate_fixtures_is_reproducible_and_seeded(run_cli, tmp_path):
    def generate(out, *extra):
        args = [
            "generate-fixtures",
            "--weeks",
            "1",
            "--end-week",
            "2025-W38",
            "--out",
            out,
            "--config",
            ROOT / "report.yaml",
        ]
        assert run_cli([*args, *extra]).exit_code == 0
        return {path.name: path.read_bytes() for path in sorted(out.glob("*.xlsx"))}

    first = generate(tmp_path / "a")
    assert generate(tmp_path / "b") == first
    other = generate(tmp_path / "c", "--seed", "1")
    assert set(other) == set(first)
    assert all(content_fingerprint(other[name]) != content_fingerprint(first[name]) for name in first)
