import io
import os
import zipfile
from datetime import UTC, datetime

import pytest
from hypothesis import given, settings
from hypothesis import strategies as st
from openpyxl import Workbook, load_workbook

from buildreport.reproducible import (
    MEMBER_ATTR,
    TimestampError,
    environment,
    normalise_zip,
    run_timestamp,
    save_workbook_bytes,
    write_atomic,
    zip_time,
)

WHEN = datetime(2025, 9, 22, 6, 0, tzinfo=UTC)
CORE = (
    b'<?xml version="1.0"?><cp:coreProperties xmlns:cp="c" xmlns:dcterms="d">'
    b'<dcterms:created xsi:type="dcterms:W3CDTF">2030-01-01T10:11:12Z</dcterms:created>'
    b'<dcterms:modified xsi:type="dcterms:W3CDTF">2031-02-03T04:05:06Z</dcterms:modified>'
    b"</cp:coreProperties>"
)


def make_zip(members):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_STORED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    return buffer.getvalue()


def test_timestamp_from_source_date_epoch():
    stamp = run_timestamp({"SOURCE_DATE_EPOCH": "1758520800"})
    assert stamp.value == WHEN
    assert stamp.source == "SOURCE_DATE_EPOCH"
    assert stamp.iso == "2025-09-22T06:00:00Z"
    assert stamp.display == "2025-09-22 06:00 UTC"
    assert stamp.naive == datetime(2025, 9, 22, 6, 0)


def test_timestamp_from_the_clock_drops_microseconds():
    stamp = run_timestamp({}, now=datetime(2025, 9, 22, 6, 0, 30, 123456, tzinfo=UTC))
    assert stamp.source == "clock"
    assert stamp.value == datetime(2025, 9, 22, 6, 0, 30, tzinfo=UTC)


@pytest.mark.parametrize("raw", ["abc", "-5", "1.5"])
def test_invalid_source_date_epoch(raw):
    with pytest.raises(TimestampError):
        run_timestamp({"SOURCE_DATE_EPOCH": raw})


def test_zip_time_is_clamped_and_even():
    assert zip_time(datetime(1970, 1, 1, tzinfo=UTC)) == (1980, 1, 1, 0, 0, 0)
    assert zip_time(datetime(2025, 9, 22, 6, 0, 1, tzinfo=UTC)) == (2025, 9, 22, 6, 0, 0)
    assert zip_time(WHEN) == (2025, 9, 22, 6, 0, 0)


def test_normalise_sets_times_attributes_and_core_properties():
    data = make_zip(
        {"[Content_Types].xml": b"<Types/>", "docProps/core.xml": CORE, "xl/workbook.xml": b"<w/>"}
    )
    result = normalise_zip(data, WHEN)
    with zipfile.ZipFile(io.BytesIO(result)) as archive:
        infos = archive.infolist()
        assert [info.filename for info in infos] == [
            "[Content_Types].xml",
            "docProps/core.xml",
            "xl/workbook.xml",
        ]
        for info in infos:
            assert info.date_time == (2025, 9, 22, 6, 0, 0)
            assert info.create_system == 0
            assert info.external_attr == MEMBER_ATTR
            assert info.compress_type == zipfile.ZIP_DEFLATED
        core = archive.read("docProps/core.xml")
    assert b">2025-09-22T06:00:00Z</dcterms:created>" in core
    assert b">2025-09-22T06:00:00Z</dcterms:modified>" in core
    assert b"2030" not in core and b"2031" not in core


@settings(max_examples=50)
@given(
    st.dictionaries(
        st.text(alphabet="abcdefghij/", min_size=1, max_size=12).filter(lambda name: not name.endswith("/")),
        st.binary(max_size=200),
        min_size=1,
        max_size=5,
    )
)
def test_normalising_twice_equals_once(members):
    once = normalise_zip(make_zip(members), WHEN)
    assert normalise_zip(once, WHEN) == once


def test_saved_workbooks_are_byte_identical_and_carry_the_timestamp():
    def workbook():
        book = Workbook()
        book.active["A1"] = "Cedarline"
        return book

    first = save_workbook_bytes(workbook(), WHEN)
    second = save_workbook_bytes(workbook(), WHEN)
    assert first == second
    reloaded = load_workbook(io.BytesIO(first))
    assert reloaded.properties.created == datetime(2025, 9, 22, 6, 0)
    assert reloaded.properties.modified == datetime(2025, 9, 22, 6, 0)


def test_write_atomic_replaces_and_leaves_no_temporary_file(tmp_path):
    target = tmp_path / "out.bin"
    write_atomic(target, b"one")
    write_atomic(target, b"two")
    assert target.read_bytes() == b"two"
    assert sorted(os.listdir(tmp_path)) == ["out.bin"]


def test_environment_lists_the_versions():
    env = environment()
    assert list(env) == ["python", "platform", "openpyxl", "pydantic", "PyYAML", "click"]
    assert env["openpyxl"].startswith("3.1.")
