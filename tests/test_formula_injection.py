"""Text taken from the input workbooks is written as text, never as a formula (final review I1)."""

import io
import re
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from buildreport.workbook import formula_cells
from buildreport.workbook.styles import put_text

from conftest import ROOT, STEM

HOSTILE_CODE = '=HYPERLINK("https://example.invalid","x")'
HOSTILE_FILE = "=1+1.xlsx"


def craft(folder, cfg):
    """A file named like a formula, and a return whose cover Branch code is formula-shaped text."""
    cover = cfg.model.template.cover
    (folder / "CL-04_2025-W38.xlsx").rename(folder / HOSTILE_FILE)
    path = folder / "CL-07_2025-W38.xlsx"
    book = load_workbook(path)
    sheet = book[cover.sheet.name]
    assert sheet["A3"].value == cover.labels.branch_code
    sheet["B3"].value = HOSTILE_CODE
    sheet["B3"].data_type = "s"  # what a typed apostrophe plus equals sign gives in Excel
    book.save(path)


@pytest.fixture()
def crafted_report(run_cli, scenario, cfg, tmp_path):
    craft(scenario, cfg)
    out = tmp_path / "out"
    result = run_cli(
        ["run", "--period", "2025-W38", "--input", scenario, "--out", out, "--config", ROOT / "report.yaml"]
    )
    assert result.exit_code == 5, result.output  # the crafted cover code is an error finding
    return (out / f"{STEM}.xlsx").read_bytes()


def cells_with(ws, text):
    return [cell for row in ws.iter_rows() for cell in row if cell.value == text]


def test_input_derived_text_stays_text_in_the_built_report(crafted_report):
    book = load_workbook(io.BytesIO(crafted_report))
    run_info, data = book["Run info"], book["Data"]

    code_cells = cells_with(run_info, HOSTILE_CODE)
    assert code_cells, "the cover text must be shown on Run info"
    file_cells = cells_with(run_info, HOSTILE_FILE)
    assert len(file_cells) >= 2  # the input list and the findings
    source_cells = cells_with(data, HOSTILE_FILE)
    assert len(source_cells) == 31  # one branch-week of records
    for cell in [*code_cells, *file_cells, *source_cells]:
        assert cell.data_type == "s", cell.coordinate
        assert isinstance(cell.value, str)

    for ws in (run_info, data):
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", (ws.title, cell.coordinate)


def test_only_the_total_formulas_remain(crafted_report, cfg, demo_result):
    book = load_workbook(io.BytesIO(crafted_report))
    found = {
        (ws.title, cell.coordinate)
        for ws in book
        for row in ws.iter_rows()
        for cell in row
        if cell.data_type == "f"
    }
    want = {(cell.sheet, cell.ref) for cell in formula_cells(cfg, demo_result)}
    assert len(want) == 30
    assert found == want


def test_no_formula_element_in_the_text_sheets(crafted_report):
    with zipfile.ZipFile(io.BytesIO(crafted_report)) as archive:
        names = [info.filename for info in archive.infolist()]
        sheets = {
            int(re.search(r"sheet(\d+)\.xml", name)[1]): name for name in names if "worksheets/sheet" in name
        }
        data_sheet, run_info_sheet = archive.read(sheets[4]), archive.read(sheets[5])
        total_sheet = archive.read(sheets[2])
    assert b"<f>" not in data_sheet and b"<f " not in data_sheet
    assert b"<f>" not in run_info_sheet and b"<f " not in run_info_sheet
    assert total_sheet.count(b"<f>") > 0  # the totals are still formulas


@pytest.mark.parametrize(
    "text",
    [
        '=HYPERLINK("https://example.invalid","x")',
        "=1+1",
        "+1+1",
        "-1-1",
        "@SUM(1)",
        "\t=1",
        "\r=1",
        "#REF!",
        "plain",
    ],
)
def test_put_text_pins_any_string_to_text(text):
    ws = Workbook().active
    cell = put_text(ws, 1, 1, text)
    assert cell.data_type == "s"
    assert cell.value == text


def test_put_text_leaves_other_values_alone():
    ws = Workbook().active
    assert put_text(ws, 1, 1, 12.5).data_type == "n"
    assert put_text(ws, 2, 1, None).value is None
