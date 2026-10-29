import json
import shutil

import pytest
from openpyxl import load_workbook

from conftest import ROOT, STEM


def none(folder):
    return folder


def branch_missing(folder):
    (folder / "CL-07_2025-W38.xlsx").unlink()
    return folder


def identical_duplicate(folder):
    shutil.copyfile(folder / "CL-04_2025-W38.xlsx", folder / "CL-04_2025-W38_copy.xlsx")
    return folder


def different_duplicates_same_time(folder):
    book = load_workbook(folder / "CL-04_2025-W38.xlsx")
    book["Sales"]["B2"] = book["Sales"]["B2"].value + 1
    book.save(folder / "CL-04_2025-W38_copy.xlsx")
    return folder


def corrupt_file(folder):
    (folder / "CL-09_2025-W38.xlsx").write_bytes(b"this is not a workbook")
    return folder


def no_file_for_period(folder):
    for path in folder.glob("*_2025-W38*.xlsx"):
        path.unlink()
    return folder


# (name, setup, exit code normal, exit code --strict); "demo" uses the committed set.
CASES = [
    ("no deviations", none, 0, 0),
    ("demo set", None, 0, 3),
    ("current-week branch missing", branch_missing, 0, 3),
    ("identical duplicate", identical_duplicate, 0, 3),
    ("different duplicates, equal Submitted at", different_duplicates_same_time, 5, 3),
    ("corrupt file", corrupt_file, 5, 3),
    ("no file for the period", no_file_for_period, 3, 3),
]


@pytest.mark.parametrize("strict", [False, True], ids=["normal", "strict"])
@pytest.mark.parametrize(("name", "setup", "normal", "strict_code"), CASES, ids=[case[0] for case in CASES])
def test_exit_codes(run_cli, scenario, demo_dir, tmp_path, name, setup, normal, strict_code, strict):
    folder = demo_dir if setup is None else setup(scenario)
    out = tmp_path / "out"
    args = ["run", "--period", "2025-W38", "--input", folder, "--out", out, "--config", ROOT / "report.yaml"]
    result = run_cli([*args, "--strict"] if strict else args)
    expected = strict_code if strict else normal
    assert result.exit_code == expected, result.output
    manifest = json.loads((out / f"{STEM}.manifest.json").read_text(encoding="utf-8"))
    assert manifest["exit_code"] == expected
    workbook = out / f"{STEM}.xlsx"
    if expected == 3:
        assert manifest["status"] == "blocked" and not workbook.exists()
    else:
        assert manifest["status"] == "built" and workbook.is_file()
    if name == "current-week branch missing" and not strict:
        assert "11 branches reported, 12 expected" in load_workbook(workbook)["Summary"]["A2"].value
    if name == "different duplicates, equal Submitted at" and not strict:
        codes = [item["code"] for item in manifest["findings"]]
        assert codes.count("DUPLICATE_AMBIGUOUS") == 2
