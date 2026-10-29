"""Seeded generator for the weekly branch workbooks (template version 3) and the demo deviation profile.

The parameters describe invented data, so they live here and not in report.yaml. Every file gets its own
random.Random(f"cedarline:{seed}:{code}:{week}") and draws only through random(), whose sequence Python keeps
stable across versions; normal draws use Box-Muller and are redrawn until they fall inside their bounds.
"""

from __future__ import annotations

import hashlib
import io
import math
import random
import tempfile
from collections.abc import Sequence
from dataclasses import dataclass, replace
from datetime import UTC, date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Literal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from buildreport.config import Config
from buildreport.isoweek import IsoWeek, window
from buildreport.reproducible import save_workbook_bytes

DEFAULT_SEED = 2207
DEMO_END = IsoWeek(2025, 38)
DEMO_WEEKS = 8
GROWTH_BASE = date(2024, 12, 30)
FORMATS = ("Flagship", "High street", "Neighbourhood", "Retail park", "Transit kiosk")


@dataclass(frozen=True)
class BranchParams:
    base: float  # net sales EUR per week
    basket: float  # average basket EUR
    labour_target: float  # labour cost as a share of net sales
    waste_factor: float


@dataclass(frozen=True)
class GroupParams:
    price: float
    waste_rate: float  # waste EUR / net sales
    mix: tuple[float, float, float, float, float]  # percent of branch sales, in FORMATS order


@dataclass(frozen=True)
class RoleParams:
    rate: float  # employer cost EUR per hour
    share: float  # share of labour cost


BRANCH_PARAMS: dict[str, BranchParams] = {
    "CL-01": BranchParams(41_500, 8.90, 0.270, 0.90),
    "CL-02": BranchParams(33_000, 8.20, 0.280, 1.00),
    "CL-03": BranchParams(30_500, 8.30, 0.285, 1.12),
    "CL-04": BranchParams(24_000, 7.60, 0.300, 0.95),
    "CL-05": BranchParams(22_500, 7.50, 0.305, 1.25),
    "CL-06": BranchParams(21_000, 7.70, 0.295, 0.85),
    "CL-07": BranchParams(15_500, 5.80, 0.310, 0.80),
    "CL-08": BranchParams(23_500, 7.40, 0.300, 1.05),
    "CL-09": BranchParams(29_000, 8.10, 0.280, 0.92),
    "CL-10": BranchParams(26_500, 9.40, 0.265, 1.18),
    "CL-11": BranchParams(19_500, 7.30, 0.325, 1.00),
    "CL-12": BranchParams(14_000, 5.60, 0.315, 0.88),
}
GROUP_PARAMS: dict[str, GroupParams] = {
    "Bread": GroupParams(3.40, 0.075, (22, 20, 27, 21, 8)),
    "Viennoiserie": GroupParams(2.60, 0.085, (17, 18, 16, 15, 24)),
    "Cakes and Pastries": GroupParams(3.90, 0.050, (19, 16, 18, 20, 10)),
    "Savoury and Sandwiches": GroupParams(5.80, 0.065, (20, 22, 17, 21, 22)),
    "Hot Drinks": GroupParams(3.50, 0.004, (16, 18, 16, 15, 30)),
    "Cold Drinks and Grocery": GroupParams(2.40, 0.006, (6, 6, 6, 8, 6)),
}
ROLE_PARAMS: dict[str, RoleParams] = {
    "Bakery": RoleParams(21.40, 0.40),
    "Counter": RoleParams(17.80, 0.48),
    "Supervisor": RoleParams(25.60, 0.12),
}
_SEASONALITY = (
    (1, 1, 0.86),
    (2, 5, 0.90),
    (6, 12, 0.96),
    (13, 22, 1.00),
    (23, 26, 1.02),
    (27, 34, 0.95),
    (35, 39, 1.00),
    (40, 47, 1.03),
    (48, 48, 1.08),
    (49, 49, 1.12),
    (50, 50, 1.18),
    (51, 51, 1.24),
    (52, 52, 1.05),
    (53, 53, 0.92),
)
_NUMBER_FORMATS = {"money": "#,##0.00", "count": "#,##0", "hours": "0.0", "text": "General"}
_TITLE = Font(name="Calibri", size=14, bold=True)
_LABEL = Font(name="Calibri", bold=True)
_HEADER = Font(name="Calibri", bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F2A36")


def seasonality(week: IsoWeek) -> float:
    for first, last, factor in _SEASONALITY:
        if first <= week.week <= last:
            return factor
    raise ValueError(f"no seasonality for {week}")


def growth(week: IsoWeek) -> float:
    """3 % a year, compounded from the Monday of 2025-W01."""
    return 1.03 ** ((week.monday - GROWTH_BASE).days / 365.25)


def mix_for(fmt: str, week: IsoWeek) -> dict[str, float]:
    """Product-group shares of branch sales for a format, adjusted for winter and summer, summing to 1."""
    index = FORMATS.index(fmt)
    weights = {group: params.mix[index] / 100 for group, params in GROUP_PARAMS.items()}
    if week.week >= 44 or week.week <= 10:
        weights["Hot Drinks"] *= 1.12
        weights["Cold Drinks and Grocery"] *= 0.85
    elif 23 <= week.week <= 34:
        weights["Hot Drinks"] *= 0.88
        weights["Cold Drinks and Grocery"] *= 1.20
    total = math.fsum(weights.values())
    return {group: weight / total for group, weight in weights.items()}


def normal(rng: random.Random, sigma: float, low: float | None = None, high: float | None = None) -> float:
    """N(1, sigma) by Box-Muller from rng.random() only, redrawn until inside [low, high] (default 3 sigma)."""
    lower = 1 - 3 * sigma if low is None else low
    upper = 1 + 3 * sigma if high is None else high
    while True:
        u1 = rng.random()
        u2 = rng.random()
        z = math.sqrt(-2.0 * math.log(1.0 - u1)) * math.cos(2.0 * math.pi * u2)
        value = 1.0 + sigma * z
        if lower <= value <= upper:
            return value


@dataclass(frozen=True)
class SalesLine:
    group: str
    net_sales_eur: float
    units_sold: int
    waste_units: int
    waste_eur: float


@dataclass(frozen=True)
class LabourLine:
    role: str
    hours: float
    cost_eur: float


@dataclass(frozen=True)
class ReturnSpec:
    code: str
    name: str
    week: IsoWeek
    transactions: int
    submitted_at: datetime  # local time, naive, as a manager would type it
    submitted_by: str
    sales: tuple[SalesLine, ...]
    labour: tuple[LabourLine, ...]
    file_week: IsoWeek | None = None  # week in the file name when it differs from the cover
    suffix: str = ""  # appended to the file name, such as "_resend"
    sales_sheet: str = "Sales"
    notes_column: bool = False
    cover_week_blank: bool = False

    @property
    def file_name(self) -> str:
        return f"{self.code}_{self.file_week or self.week}{self.suffix}.xlsx"


def make_spec(cfg: Config, code: str, week: IsoWeek, seed: int = DEFAULT_SEED) -> ReturnSpec:
    """One branch return by the rules of spec section 4.4."""
    branch = cfg.branch(code)
    if branch is None or code not in BRANCH_PARAMS:
        raise ValueError(f"the generator has no parameters for branch {code}")
    params = BRANCH_PARAMS[code]
    rng = random.Random(f"cedarline:{seed}:{code}:{week}")
    season = seasonality(week)
    trend = growth(week)
    branch_sales = params.base * season * trend * normal(rng, 0.035, 0.90, 1.10)
    mix = mix_for(branch.format, week)
    sales: list[SalesLine] = []
    for group in cfg.product_groups:
        group_params = GROUP_PARAMS[group]
        net = round(branch_sales * mix[group] * normal(rng, 0.05, 0.85, 1.15), 2)
        units = round(net / group_params.price * normal(rng, 0.02))
        waste_eur = round(
            net * group_params.waste_rate * params.waste_factor * normal(rng, 0.12, 0.7, 1.3), 2
        )
        sales.append(SalesLine(group, net, units, round(waste_eur / group_params.price), waste_eur))
    total = math.fsum(line.net_sales_eur for line in sales)
    transactions = round(total / params.basket * normal(rng, 0.03))
    labour_cost = params.labour_target * params.base * trend * (0.7 + 0.3 * season) * normal(rng, 0.015)
    labour: list[LabourLine] = []
    for role in cfg.roles:
        role_params = ROLE_PARAMS[role]
        hours = round(labour_cost * role_params.share / role_params.rate * 2) / 2
        labour.append(LabourLine(role, hours, round(hours * role_params.rate, 2)))
    number = int(code.split("-")[1])
    submitted = datetime.combine(week.sunday + timedelta(days=1), time(8, 0)) + timedelta(minutes=7 * number)
    return ReturnSpec(
        code=code,
        name=branch.name,
        week=week,
        transactions=transactions,
        submitted_at=submitted,
        submitted_by=f"{branch.name.lower()}.manager@cedarline.example",
        sales=tuple(sales),
        labour=tuple(labour),
    )


def apply_demo(specs: list[ReturnSpec], end: IsoWeek) -> list[ReturnSpec]:
    """The six deviations of the demo profile relative to the end week E (spec section 4.3).

    A deviation whose week is not among the generated weeks is skipped, because its spec is absent.
    """
    late = end.shift(-10)
    names = {spec.file_name for spec in specs}
    result: list[ReturnSpec] = []
    for spec in specs:
        if spec.week == end and spec.code == "CL-03":
            spec = replace(spec, sales_sheet=f"Sales wk{end.week:02d}")
        elif spec.week == end and spec.code == "CL-05":
            spec = replace(spec, notes_column=True)
        elif spec.week == end and spec.code == "CL-08":
            spec = replace(spec, cover_week_blank=True)
        elif spec.week == end and spec.code == "CL-11":
            suffix = "_late" if f"CL-11_{late}.xlsx" in names else ""
            spec = replace(spec, file_week=late, suffix=suffix)
        elif spec.week == end.shift(-2) and spec.code == "CL-12":
            continue
        elif spec.week == end.shift(-1) and spec.code == "CL-10":
            monday = spec.week.sunday + timedelta(days=1)
            result.append(replace(spec, submitted_at=datetime.combine(monday, time(8, 40))))
            corrected = tuple(
                replace(line, net_sales_eur=round(line.net_sales_eur * 1.035, 2))
                if line.group == "Bread"
                else line
                for line in spec.sales
            )
            spec = replace(
                spec, submitted_at=datetime.combine(monday, time(14, 12)), sales=corrected, suffix="_resend"
            )
        result.append(spec)
    return result


def _table(sheet: Any, headers: list[str], rows: list[list[Any]], types: list[str]) -> None:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = _HEADER
        cell.fill = _HEADER_FILL
        sheet.column_dimensions[cell.column_letter].width = 24 if column == 1 else 15
    for row_number, row in enumerate(rows, start=2):
        for column, (value, kind) in enumerate(zip(row, types, strict=True), start=1):
            cell = sheet.cell(row=row_number, column=column, value=value)
            cell.number_format = _NUMBER_FORMATS[kind]


def build_workbook(spec: ReturnSpec, cfg: Config) -> Workbook:
    """The branch workbook of template version 3 for one spec, with its deviations."""
    template = cfg.model.template
    labels = template.cover.labels
    book = Workbook()
    book.properties.creator = "Cedarline Bakeries"
    book.properties.title = f"Weekly branch return {spec.code} {spec.week}"
    cover = book.active
    cover.title = template.cover.sheet.name
    cover["A1"] = "Cedarline Bakeries · weekly branch return"
    cover["A1"].font = _TITLE
    entries: list[tuple[str, Any]] = [
        (labels.branch_code, spec.code),
        (labels.branch_name, spec.name),
        (labels.week, None if spec.cover_week_blank else str(spec.week)),
        (labels.transactions, spec.transactions),
        (labels.submitted_at, spec.submitted_at),
        (labels.submitted_by, spec.submitted_by),
        (labels.template_version, template.version),
    ]
    for row, (label, value) in enumerate(entries, start=3):
        cover.cell(row=row, column=1, value=label).font = _LABEL
        cell = cover.cell(row=row, column=2, value=value)
        if isinstance(value, datetime):
            cell.number_format = "yyyy-mm-dd hh:mm"
    cover["A11"] = "Weekly totals only, Monday to Sunday. Send to head office every Monday by 12:00."
    cover.column_dimensions["A"].width = 18
    cover.column_dimensions["B"].width = 36

    columns = template.sales.columns
    headers = [template.sales.key.header, *(column.header for column in columns.values())]
    types = ["text", *(column.type for column in columns.values())]
    rows: list[list[Any]] = [
        [line.group, line.net_sales_eur, line.units_sold, line.waste_units, line.waste_eur]
        for line in spec.sales
    ]
    if spec.notes_column:
        position = list(columns).index("units_sold") + 2
        headers.insert(position, "Notes")
        types.insert(position, "text")
        for index, row in enumerate(rows):
            row.insert(position, "Oven 2 serviced on Tuesday" if index == 0 else None)
    _table(book.create_sheet(spec.sales_sheet), headers, rows, types)

    labour_columns = template.labour.columns
    _table(
        book.create_sheet(template.labour.sheet.name),
        [template.labour.key.header, *(column.header for column in labour_columns.values())],
        [[line.role, line.hours, line.cost_eur] for line in spec.labour],
        ["text", *(column.type for column in labour_columns.values())],
    )
    return book


def to_bytes(spec: ReturnSpec, cfg: Config) -> bytes:
    """Saved reproducibly with Submitted at as the timestamp."""
    return save_workbook_bytes(build_workbook(spec, cfg), spec.submitted_at.replace(tzinfo=UTC))


@dataclass(frozen=True)
class FixtureFile:
    name: str
    data: bytes


def generate(
    cfg: Config,
    end: IsoWeek,
    weeks: int,
    deviations: Literal["none", "demo"] = "demo",
    seed: int = DEFAULT_SEED,
) -> list[FixtureFile]:
    """``weeks`` weeks of twelve files ending with ``end``, sorted by file name."""
    if not 1 <= weeks <= 104:
        raise ValueError("weeks must be between 1 and 104")
    specs = [make_spec(cfg, code, week, seed) for week in window(end, weeks) for code in cfg.branch_codes]
    if deviations == "demo":
        specs = apply_demo(specs, end)
    specs.sort(key=lambda spec: spec.file_name)
    return [FixtureFile(spec.file_name, to_bytes(spec, cfg)) for spec in specs]


def write_fixtures(files: Sequence[FixtureFile], out: Path, force: bool = False) -> None:
    """Write the files; without ``force`` nothing is written when any target exists."""
    out.mkdir(parents=True, exist_ok=True)
    existing = [item.name for item in files if (out / item.name).exists()]
    if existing and not force:
        raise FileExistsError(
            f"{len(existing)} files already exist in {out} (first: {existing[0]}); use --force to overwrite"
        )
    for item in files:
        (out / item.name).write_bytes(item.data)


def content_fingerprint(data: bytes) -> str:
    """SHA-256 over sheet names and typed cell values; equal on every platform when the content is equal."""
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    digest = hashlib.sha256()
    try:
        for sheet in book.worksheets:
            digest.update(f"sheet:{sheet.title}\n".encode())
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                for column, value in enumerate(row, start=1):
                    if value is not None:
                        digest.update(f"{row_number}:{column}={type(value).__name__}:{value!r}\n".encode())
    finally:
        book.close()
    return digest.hexdigest()


def fingerprint_folder(folder: Path) -> dict[str, str]:
    return {path.name: content_fingerprint(path.read_bytes()) for path in sorted(folder.glob("*.xlsx"))}


def regenerate_fingerprints(cfg: Config) -> dict[str, str]:
    """The demo set regenerated in a temporary folder, as fingerprints by file name."""
    with tempfile.TemporaryDirectory(prefix="cedarline-fixtures-") as folder:
        write_fixtures(
            generate(cfg, DEMO_END, DEMO_WEEKS, deviations="demo", seed=DEFAULT_SEED), Path(folder)
        )
        return fingerprint_folder(Path(folder))
