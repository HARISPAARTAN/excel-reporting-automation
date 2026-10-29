"""Template reader: one branch workbook (template version 3) into a BranchReturn plus findings.

Workbooks are opened with data_only=True; numbers are never parsed from strings.
"""

from __future__ import annotations

import fnmatch
import hashlib
import io
import math
from dataclasses import dataclass
from datetime import date, datetime, time
from typing import Any

from openpyxl import load_workbook

from buildreport.config import Config, SheetSpec, TableSpec, normalise
from buildreport.discover import SourceFile
from buildreport.findings import Finding, finding
from buildreport.isoweek import IsoWeek

COVER_ROWS = 30


@dataclass(frozen=True)
class BranchReturn:
    file: str
    branch_code: str
    transactions: int
    submitted_at: datetime
    sales: dict[str, dict[str, float]]  # product group -> measure -> value
    labour: dict[str, dict[str, float]]  # role -> measure -> value


@dataclass(frozen=True)
class ReadResult:
    source: SourceFile
    size: int
    sha256: str
    cover_read: bool
    cover_branch: str | None
    cover_week_text: str | None
    cover_week: IsoWeek | None
    submitted_at: datetime | None
    ret: BranchReturn | None  # None when an error finding excludes the file
    findings: tuple[Finding, ...]

    @property
    def has_error(self) -> bool:
        return any(item.severity == "error" for item in self.findings)


class _Collector:
    def __init__(self, file: str) -> None:
        self.file = file
        self.items: list[Finding] = []

    def add(self, code: str, message: str) -> None:
        self.items.append(finding(code, self.file, message))

    @property
    def failed(self) -> bool:
        return any(item.severity == "error" for item in self.items)


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _number(value: Any, kind: str) -> tuple[float | None, str | None]:
    """(number, problem) for a cell of a numeric column."""
    if _blank(value):
        return None, "blank"
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        if isinstance(value, str):
            return None, f"text {value!r} in a numeric cell"
        return None, f"{type(value).__name__} in a numeric cell"
    number = float(value)
    if not math.isfinite(number):
        return None, "not a finite number"
    if number < 0:
        return None, f"negative value {value}"
    if kind == "count" and not number.is_integer():
        return None, f"{value} is not a whole number"
    return number, None


def find_sheet(names: list[str], spec: SheetSpec) -> tuple[str | None, str | None, str | None]:
    """(sheet name, alias used, problem): the exact name first (case-insensitive), then the glob aliases."""
    wanted = normalise(spec.name)
    for name in names:
        if normalise(name) == wanted:
            return name, None, None
    matches: list[tuple[str, str]] = []
    for name in names:
        folded = normalise(name)
        alias = next((item for item in spec.aliases if fnmatch.fnmatchcase(folded, normalise(item))), None)
        if alias is not None:
            matches.append((name, alias))
    if len(matches) == 1:
        return matches[0][0], matches[0][1], None
    if not matches:
        aliases = ", ".join(spec.aliases) or "no alias"
        return None, None, f"no sheet named {spec.name} or matching {aliases}"
    listed = ", ".join(repr(name) for name, _ in matches)
    return None, None, f"{len(matches)} sheets match {spec.name}: {listed}"


def _locate(book: Any, spec: SheetSpec, out: _Collector) -> Any:
    name, alias, problem = find_sheet(list(book.sheetnames), spec)
    if name is None:
        out.add("SHEET_MISSING", problem or f"no sheet {spec.name}")
        return None
    if alias is not None:
        out.add("SHEET_RENAMED", f"sheet {name!r} used for {spec.name} (alias {alias!r})")
    return book[name]


def _cover_values(sheet: Any, labels: dict[str, str]) -> dict[str, Any]:
    """Column B next to the labels found in A1:A30, matched case-insensitively; the first match wins."""
    wanted = {normalise(text): key for key, text in labels.items()}
    found: dict[str, Any] = {}
    for label, value in sheet.iter_rows(
        min_row=1, max_row=COVER_ROWS, min_col=1, max_col=2, values_only=True
    ):
        if isinstance(label, str):
            key = wanted.get(normalise(label))
            if key is not None and key not in found:
                found[key] = value
    return found


def _read_table(
    sheet: Any, spec: TableSpec, allowed: tuple[str, ...], what: str, out: _Collector
) -> dict[str, dict[str, float]]:
    rows = list(sheet.iter_rows(values_only=True))
    header: tuple[Any, ...] = rows[0] if rows else ()
    positions: dict[str, int] = {}
    for index, value in enumerate(header):
        if _blank(value):
            continue
        key = normalise(str(value))
        if key in positions:
            out.add("VALUE_INVALID", f"column {str(value)!r} appears twice on sheet {sheet.title}")
            continue
        positions[key] = index
    key_header = normalise(spec.key.header)
    wanted = {normalise(column.header): name for name, column in spec.columns.items()}
    complete = True
    for header_text in (spec.key.header, *(column.header for column in spec.columns.values())):
        if normalise(header_text) not in positions:
            out.add("COLUMN_MISSING", f"no column {header_text!r} on sheet {sheet.title}")
            complete = False
    reported: set[str] = set()
    for value in header:
        if _blank(value):
            continue
        folded = normalise(str(value))
        if folded != key_header and folded not in wanted and folded not in reported:
            reported.add(folded)
            out.add("EXTRA_COLUMN", f"column {str(value)!r} on sheet {sheet.title} ignored")
    if not complete:
        return {}

    canonical = {normalise(item): item for item in allowed}
    numeric = {name: column.type for name, column in spec.columns.items() if column.type != "text"}
    values: dict[str, dict[str, float]] = {}
    for row_number, row in enumerate(rows[1:], start=2):
        if all(_blank(cell) for cell in row):
            continue
        key_value = row[positions[key_header]]
        if _blank(key_value):
            out.add(
                "VALUE_INVALID",
                f"row {row_number} on sheet {sheet.title} has values but no {spec.key.header}",
            )
            continue
        label = str(key_value).strip()
        folded = normalise(label)
        if folded == "total":
            continue
        if folded not in canonical:
            out.add("ROW_UNKNOWN", f"{what} {label!r} on sheet {sheet.title} is not in report.yaml")
            continue
        name = canonical[folded]
        if name in values:
            out.add("VALUE_INVALID", f"{what} {name!r} appears twice on sheet {sheet.title}")
            continue
        entry: dict[str, float] = {}
        for measure, kind in numeric.items():
            column = spec.columns[measure]
            number, problem = _number(row[positions[normalise(column.header)]], kind)
            if problem is not None or number is None:
                out.add(
                    "VALUE_INVALID",
                    f"{column.header} for {name} on sheet {sheet.title}: {problem or 'blank'}",
                )
            else:
                entry[measure] = number
        values[name] = entry
    for name in allowed:
        if name not in values:
            out.add("ROW_MISSING", f"{what} {name!r} missing on sheet {sheet.title}; counted as zero")
            values[name] = {measure: 0.0 for measure in numeric}
    return {name: values[name] for name in allowed}


def read_return(source: SourceFile, cfg: Config) -> ReadResult:
    """Read one workbook; every problem becomes a finding, an error finding leaves ``ret`` as None."""
    data = source.path.read_bytes()
    size, digest = len(data), hashlib.sha256(data).hexdigest()
    out = _Collector(source.name)
    if not source.matches:
        out.add(
            "FILE_NAME", "file name does not match input.file_pattern; branch and week taken from the cover"
        )
    elif source.week is None:
        out.add(
            "FILE_NAME",
            f"week {source.week_text!r} in the file name is not an ISO week; week taken from the cover",
        )
    try:
        book = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # openpyxl raises many types for corrupt, encrypted or foreign files
        out.add("FILE_UNREADABLE", f"cannot open the workbook ({type(exc).__name__})")
        return ReadResult(source, size, digest, False, None, None, None, None, None, tuple(out.items))

    template = cfg.model.template
    labels = template.cover.labels
    cover_branch: str | None = None
    cover_week_text: str | None = None
    cover_week: IsoWeek | None = None
    submitted: datetime | None = None
    transactions: int | None = None
    cover = _locate(book, template.cover.sheet, out)
    if cover is not None:
        values = _cover_values(cover, labels.model_dump())
        raw_branch = values.get("branch_code")
        if _blank(raw_branch):
            out.add("BRANCH_INVALID", f"{labels.branch_code} on the cover is blank")
        else:
            cover_branch = str(raw_branch).strip()
            if cfg.branch(cover_branch) is None:
                out.add("BRANCH_INVALID", f"branch code {cover_branch!r} is not in report.yaml")
            elif source.matches and source.branch != cover_branch:
                out.add("BRANCH_INVALID", f"cover says {cover_branch}, the file name says {source.branch}")
        raw_week = values.get("week")
        if not _blank(raw_week):
            cover_week_text = str(raw_week).strip()
            try:
                cover_week = IsoWeek.parse(cover_week_text)
            except ValueError:
                cover_week = None
        number, problem = _number(values.get("transactions"), "count")
        if problem is not None or number is None:
            out.add("VALUE_INVALID", f"{labels.transactions} on the cover: {problem or 'blank'}")
        else:
            transactions = int(number)
        raw_submitted = values.get("submitted_at")
        if isinstance(raw_submitted, datetime):
            submitted = raw_submitted
        elif isinstance(raw_submitted, date):
            submitted = datetime.combine(raw_submitted, time())
        else:
            reason = "blank" if _blank(raw_submitted) else "not a date and time"
            out.add("VALUE_INVALID", f"{labels.submitted_at} on the cover: {reason}")
        version = values.get("template_version")
        if isinstance(version, bool) or version != template.version:
            shown = "blank" if _blank(version) else repr(version)
            out.add("TEMPLATE_VERSION", f"template version {shown}, read as {template.version}")

    sales: dict[str, dict[str, float]] = {}
    labour: dict[str, dict[str, float]] = {}
    sales_sheet = _locate(book, template.sales.sheet, out)
    if sales_sheet is not None:
        sales = _read_table(sales_sheet, template.sales, cfg.product_groups, "product group", out)
    labour_sheet = _locate(book, template.labour.sheet, out)
    if labour_sheet is not None:
        labour = _read_table(labour_sheet, template.labour, cfg.roles, "role", out)

    ret: BranchReturn | None = None
    if not out.failed and cover_branch is not None and transactions is not None and submitted is not None:
        ret = BranchReturn(source.name, cover_branch, transactions, submitted, sales, labour)
    return ReadResult(
        source,
        size,
        digest,
        cover is not None,
        cover_branch,
        cover_week_text,
        cover_week,
        submitted,
        ret,
        tuple(out.items),
    )
