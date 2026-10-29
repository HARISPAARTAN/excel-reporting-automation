"""Shared workbook styles, number formats, page setup, conditional formatting rules and FormulaCell."""

from __future__ import annotations

from collections.abc import Callable, Iterable, Mapping
from dataclasses import dataclass
from typing import Any

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from buildreport.expr import ExprError, Node, Sum, to_excel

FONT = "Calibri"
HEADER_FILL = "1F2A36"
TITLE_COLOUR = "14181D"
NOTE_COLOUR = "5B6472"
BAD_FILL = "FDE2E1"
BAD_TEXT = "9B1C1C"
WARN_FILL = "FEF3C7"
WARN_TEXT = "92400E"
POSITIVE = "15803D"
NEGATIVE = "B91C1C"
ACCENT = "F2A93B"

NUMBER_FORMATS = {
    "eur": '"€"#,##0',
    "eur2": '"€"#,##0.00',
    "count": "#,##0",
    "hours": "#,##0.0",
    "pct1": "0.0%",
}
CHANGE_FORMATS = {"relative": "+0.0%;-0.0%;0.0%", "points": '+0.0" pp";-0.0" pp";0.0" pp"'}
DATE_FORMAT = "yyyy-mm-dd"
VALUE_FORMAT = "#,##0.00"
RANK_CHANGE_FORMAT = "+0;-0;0"

TITLE_FONT = Font(name=FONT, size=16, bold=True, color=TITLE_COLOUR)
NOTE_FONT = Font(name=FONT, size=11, color=NOTE_COLOUR)
BOLD_FONT = Font(name=FONT, size=11, bold=True)
HEADER_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
HEADER_FILL_STYLE = PatternFill("solid", fgColor=HEADER_FILL)
TOTAL_BORDER = Border(top=Side(style="thin", color=TITLE_COLOUR))
OP_NAMES = {">": "greaterThan", ">=": "greaterThanOrEqual", "<": "lessThan", "<=": "lessThanOrEqual"}


@dataclass(frozen=True)
class FormulaCell:
    sheet: str
    ref: str
    formula: str  # with the leading "="
    expected: float | None  # the engine value the recalculated formula must equal


def footer(period: object) -> str:
    return f"Cedarline Bakeries · weekly report {period} · page &P of &N"


def bad_fill() -> PatternFill:
    return PatternFill(start_color=BAD_FILL, end_color=BAD_FILL, fill_type="solid")


def warn_fill() -> PatternFill:
    return PatternFill(start_color=WARN_FILL, end_color=WARN_FILL, fill_type="solid")


def title(ws: Any, ref: str, text: str) -> None:
    ws[ref] = text
    ws[ref].font = TITLE_FONT


def note(ws: Any, ref: str, text: str) -> None:
    ws[ref] = text
    ws[ref].font = NOTE_FONT


def header(ws: Any, row: int, first_column: int, labels: Iterable[str]) -> None:
    for offset, label in enumerate(labels):
        cell = ws.cell(row=row, column=first_column + offset, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL_STYLE
        cell.alignment = Alignment(
            horizontal="left" if offset == 0 else "center", vertical="center", wrap_text=True
        )


def put(ws: Any, row: int, column: int, value: Any, number_format: str | None = None) -> Any:
    cell = ws.cell(row=row, column=column, value=value)
    if number_format is not None:
        cell.number_format = number_format
    return cell


def total_row(ws: Any, row: int, first_column: int, last_column: int) -> None:
    for column in range(first_column, last_column + 1):
        cell = ws.cell(row=row, column=column)
        cell.font = BOLD_FONT
        cell.border = TOTAL_BORDER


def widths(ws: Any, mapping: Mapping[str, float]) -> None:
    for column, width in mapping.items():
        ws.column_dimensions[column].width = width


def page(
    ws: Any,
    period: object,
    *,
    landscape: bool,
    fit_height: int,
    print_area: str | None = None,
    title_rows: str | None = None,
    gridlines: bool = False,
    centred: bool = False,
) -> None:
    """A4, fit to one page wide (and fit_height pages high, 0 for automatic), footer, gridlines."""
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_height
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = footer(period)
    ws.sheet_view.showGridLines = gridlines
    if print_area is not None:
        ws.print_area = print_area
    if title_rows is not None:
        ws.print_title_rows = title_rows
    if centred:
        ws.print_options.horizontalCentered = True
        ws.print_options.verticalCentered = True


def _number(value: float) -> str:
    text = repr(float(value))
    return text[:-2] if text.endswith(".0") else text


def threshold_rule(op: str, value: float, style: str) -> CellIsRule:
    fill, text = (bad_fill(), BAD_TEXT) if style == "bad" else (warn_fill(), WARN_TEXT)
    return CellIsRule(operator=OP_NAMES[op], formula=[_number(value)], fill=fill, font=Font(color=text))


def change_rules(better: str) -> list[CellIsRule]:
    """Colours for a change column: favourable green, unfavourable red; none for neutral metrics."""
    if better == "neutral":
        return []
    good, bad = ("greaterThan", "lessThan") if better == "higher" else ("lessThan", "greaterThan")
    return [
        CellIsRule(operator=good, formula=["0"], font=Font(color=POSITIVE)),
        CellIsRule(operator=bad, formula=["0"], font=Font(color=NEGATIVE)),
    ]


def _no_sum(measure: str) -> str:
    raise KeyError(measure)


def metric_formula(node: Node, column: str, first: int, last: int, ref: Callable[[str], str]) -> str | None:
    """=SUM over the column for a plain sum metric, else the expression over the total row's cells.

    None when the expression needs a cell the sheet does not have; the caller then writes no formula.
    """
    if isinstance(node, Sum):
        return f"=SUM({column}{first}:{column}{last})"
    try:
        return "=" + to_excel(node, ref, _no_sum)
    except (KeyError, ExprError):
        return None
