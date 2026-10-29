"""Sheet Summary: title, key figures, ranking, data notes; the charts are added by charts.py."""

from __future__ import annotations

from dataclasses import dataclass
from typing import TYPE_CHECKING, Any

from openpyxl.styles import Alignment

from buildreport import __version__
from buildreport.config import Config
from buildreport.findings import counts
from buildreport.isoweek import format_range
from buildreport.metrics import count_text, threshold_text
from buildreport.workbook import styles

if TYPE_CHECKING:
    from buildreport.workbook import BuildContext

SHEET = "Summary"
KPI_HEADERS = ("Metric", "This week", "Last week", "Change")
RANK_HEADERS = ("Rank", "Branch", "Revenue", "WoW", "Waste %", "Labour %")
NOTES_ROW = 13
PRINT_AREA = "A1:L36"
WIDTHS = {
    "A": 2,
    "B": 26,
    "C": 14,
    "D": 14,
    "E": 12,
    "F": 3,
    "G": 7,
    "H": 18,
    "I": 14,
    "J": 11,
    "K": 10,
    "L": 10,
}
UNRANKED = "–"


@dataclass(frozen=True)
class Layout:
    kpi_first: int
    kpi_last: int
    rank_first: int
    rank_last: int


def layout(cfg: Config) -> Layout:
    return Layout(5, 4 + len(cfg.model.summary_kpis), 5, 4 + len(cfg.branches))


def subtitle(ctx: BuildContext) -> str:
    summary = ctx.summary
    return (
        f"{format_range(summary.period)} · {summary.reporting_now} branches reported, {summary.expected} expected · "
        f"generated {ctx.timestamp.display} · build-report {__version__}"
    )


def notes(ctx: BuildContext) -> list[str]:
    """The data notes B14 to B17 (B17 only when the two weeks differ in branch count)."""
    summary = ctx.summary
    tally = counts(ctx.resolution.findings)
    lines = [
        f"Branches reported: {summary.reporting_now} this week, "
        f"{summary.reporting_previous} last week, {summary.expected} expected.",
        f"{count_text(tally['error'], 'error')}, {count_text(tally['warning'], 'warning')}; see sheet Run info.",
        "; ".join(threshold_text(hit) for hit in summary.above_threshold) or "No thresholds configured.",
    ]
    if summary.reporting_now != summary.reporting_previous:
        lines.append(
            f"Week-on-week figures compare {summary.reporting_now} branches with {summary.reporting_previous}."
        )
    return lines


def write(ws: Any, ctx: BuildContext) -> None:
    cfg, summary = ctx.cfg, ctx.summary
    lay = layout(cfg)
    styles.title(ws, "A1", f"{cfg.model.client.name} · Weekly report {summary.period}")
    styles.note(ws, "A2", subtitle(ctx))

    styles.header(ws, 4, 2, KPI_HEADERS)
    for row, kpi in enumerate(summary.kpis, start=lay.kpi_first):
        number_format = styles.NUMBER_FORMATS[kpi.format]
        ws.cell(row=row, column=2, value=kpi.label)
        styles.put(ws, row, 3, kpi.value, number_format)
        styles.put(ws, row, 4, kpi.previous, number_format)
        styles.put(ws, row, 5, kpi.change, styles.CHANGE_FORMATS[kpi.change_type])
        threshold = cfg.threshold(kpi.key)
        if threshold is not None:
            ws.conditional_formatting.add(
                f"C{row}", styles.threshold_rule(threshold.op, threshold.value, threshold.style)
            )
        for rule in styles.change_rules(kpi.better):
            ws.conditional_formatting.add(f"E{row}", rule)

    styles.header(ws, 4, 7, RANK_HEADERS)
    revenue = cfg.metric("revenue")
    for row, entry in enumerate(summary.ranking, start=lay.rank_first):
        rank_cell = ws.cell(row=row, column=7, value=entry.rank if entry.rank is not None else UNRANKED)
        rank_cell.alignment = Alignment(horizontal="right")
        ws.cell(row=row, column=8, value=entry.name)
        styles.put(ws, row, 9, entry.revenue, styles.NUMBER_FORMATS[revenue.format])
        styles.put(ws, row, 10, entry.change, styles.CHANGE_FORMATS[revenue.change])
        styles.put(ws, row, 11, entry.waste_pct, styles.NUMBER_FORMATS[cfg.metric("waste_pct").format])
        styles.put(ws, row, 12, entry.labour_pct, styles.NUMBER_FORMATS[cfg.metric("labour_pct").format])
    for rule in styles.change_rules(revenue.better):
        ws.conditional_formatting.add(f"J{lay.rank_first}:J{lay.rank_last}", rule)
    for column, metric in (("K", "waste_pct"), ("L", "labour_pct")):
        threshold = cfg.threshold(metric)
        if threshold is not None:
            ws.conditional_formatting.add(
                f"{column}{lay.rank_first}:{column}{lay.rank_last}",
                styles.threshold_rule(threshold.op, threshold.value, threshold.style),
            )

    ws.cell(row=NOTES_ROW, column=2, value="Data notes").font = styles.BOLD_FONT
    for offset, text in enumerate(notes(ctx), start=1):
        ws.cell(row=NOTES_ROW + offset, column=2, value=text)
    thresholds_row = NOTES_ROW + 3
    ws.merge_cells(start_row=thresholds_row, start_column=2, end_row=thresholds_row, end_column=5)
    ws.cell(row=thresholds_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[thresholds_row].height = 30
    ws.cell(row=NOTES_ROW + 5, column=2, value="Charts").font = styles.BOLD_FONT

    styles.widths(ws, WIDTHS)
    ws.row_dimensions[4].height = 20
    ws.sheet_properties.tabColor = styles.ACCENT
    styles.page(ws, summary.period, landscape=True, fit_height=1, print_area=PRINT_AREA, centred=True)
