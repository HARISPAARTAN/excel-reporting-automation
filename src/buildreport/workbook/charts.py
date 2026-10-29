"""The three native charts on Summary (spec 6.4): absolute references, titles, delete=False on every axis."""

from __future__ import annotations

from typing import Any

from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.utils import quote_sheetname

from buildreport.config import Config
from buildreport.workbook import branches, products, styles, summary

TITLES = {
    "bar": "Revenue by branch, this week",
    "line": "Chain revenue, last 8 weeks",
    "doughnut": "Revenue mix by product group",
}
ANCHORS = {"bar": "B19", "line": "G19", "doughnut": "G28"}
SIZES = {"bar": (12.5, 9.0), "line": (12.5, 4.4), "doughnut": (12.5, 4.4)}
EUR_AXIS = '"€"#,##0'


def _categories(ws: Any, column: str, first: int, last: int) -> AxDataSource:
    """Text categories as a string reference (numRef would make some readers show 1, 2, 3)."""
    return AxDataSource(strRef=StrRef(f=f"{quote_sheetname(ws.title)}!${column}${first}:${column}${last}"))


def _size(chart: Any, key: str) -> None:
    chart.width, chart.height = SIZES[key]


def revenue_by_branch(summary_ws: Any, cfg: Config) -> BarChart:
    lay = summary.layout(cfg)
    chart = BarChart()
    chart.type = "bar"
    chart.grouping = "clustered"
    chart.title = TITLES["bar"]
    chart.legend = None
    chart.add_data(
        Reference(summary_ws, min_col=9, min_row=lay.rank_first, max_row=lay.rank_last),
        titles_from_data=False,
    )
    series = chart.series[0]
    series.cat = _categories(summary_ws, "H", lay.rank_first, lay.rank_last)
    series.graphicalProperties = GraphicalProperties(solidFill=styles.HEADER_FILL)
    chart.x_axis.scaling.orientation = "maxMin"
    chart.y_axis.number_format = EUR_AXIS
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    _size(chart, "bar")
    return chart


def chain_trend(branches_ws: Any, cfg: Config) -> LineChart:
    lay = branches.layout(cfg)
    chart = LineChart()
    chart.title = TITLES["line"]
    chart.legend = None
    chart.add_data(
        Reference(branches_ws, min_col=3, min_row=lay.trend_first, max_row=lay.trend_last),
        titles_from_data=False,
    )
    series = chart.series[0]
    series.cat = _categories(branches_ws, "A", lay.trend_first, lay.trend_last)
    series.marker = Marker(symbol="circle", size=6)
    series.marker.graphicalProperties = GraphicalProperties(solidFill=styles.ACCENT)
    series.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=styles.ACCENT, w=28575))
    series.smooth = False
    chart.y_axis.number_format = EUR_AXIS
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    _size(chart, "line")
    return chart


def revenue_mix(products_ws: Any, cfg: Config) -> DoughnutChart:
    lay = products.layout(cfg)
    chart = DoughnutChart(holeSize=55)
    chart.title = TITLES["doughnut"]
    chart.add_data(
        Reference(products_ws, min_col=2, min_row=lay.first, max_row=lay.last), titles_from_data=False
    )
    chart.series[0].cat = _categories(products_ws, "A", lay.first, lay.last)
    chart.legend.position = "r"
    chart.dataLabels = DataLabelList(
        showPercent=True, showVal=False, showCatName=False, showSerName=False, showLegendKey=False
    )
    _size(chart, "doughnut")
    return chart


def add_charts(book: Any, cfg: Config) -> None:
    summary_ws = book[summary.SHEET]
    summary_ws.add_chart(revenue_by_branch(summary_ws, cfg), ANCHORS["bar"])
    summary_ws.add_chart(chain_trend(book[branches.SHEET], cfg), ANCHORS["line"])
    summary_ws.add_chart(revenue_mix(book[products.SHEET], cfg), ANCHORS["doughnut"])
