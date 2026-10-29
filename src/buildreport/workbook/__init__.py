"""The weekly report workbook: Summary, Branches, Products, Data and Run info (spec 6.4)."""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl import Workbook

from buildreport import __version__
from buildreport.config import Config
from buildreport.metrics import MetricsResult, Summary
from buildreport.reproducible import RunTimestamp, save_workbook_bytes
from buildreport.resolve import Record, Resolution
from buildreport.workbook import branches, charts, data, products, runinfo, summary
from buildreport.workbook.styles import FormulaCell

SHEETS = (summary.SHEET, branches.SHEET, products.SHEET, data.SHEET, runinfo.SHEET)


@dataclass(frozen=True)
class BuildContext:
    cfg: Config
    result: MetricsResult
    summary: Summary
    resolution: Resolution
    records: list[Record]
    timestamp: RunTimestamp
    strict: bool
    input_dir: str


def build(ctx: BuildContext) -> Workbook:
    """All five sheets with titles first (chart references need them), then content, then charts."""
    book = Workbook()
    first = book.active
    first.title = SHEETS[0]
    for name in SHEETS[1:]:
        book.create_sheet(name)
    summary.write(book[summary.SHEET], ctx)
    branches.write(book[branches.SHEET], ctx)
    products.write(book[products.SHEET], ctx)
    data.write(book[data.SHEET], ctx)
    runinfo.write(book[runinfo.SHEET], ctx)
    charts.add_charts(book, ctx.cfg)
    book.active = 0
    book.properties.title = f"Cedarline weekly report {ctx.result.period}"
    book.properties.creator = f"build-report {__version__}"
    return book


def formula_cells(cfg: Config, result: MetricsResult) -> list[FormulaCell]:
    """Every formula in the workbook with the engine value it must recalculate to (30 with the demo config)."""
    return [*branches.total_cells(cfg, result), *products.total_cells(cfg, result)]


def render_bytes(ctx: BuildContext) -> bytes:
    return save_workbook_bytes(build(ctx), ctx.timestamp.value)
