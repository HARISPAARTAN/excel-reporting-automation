"""Sheet Branches: one row per branch with this week's figures, a total row, and the weekly trend."""

from __future__ import annotations

from dataclasses import dataclass
from typing import TYPE_CHECKING, Any

from openpyxl.formatting.rule import DataBarRule
from openpyxl.utils import column_index_from_string

from buildreport.config import Config
from buildreport.expr import change_excel
from buildreport.metrics import MetricsResult, change
from buildreport.workbook import styles
from buildreport.workbook.styles import FormulaCell

if TYPE_CHECKING:
    from buildreport.workbook import BuildContext

SHEET = "Branches"
HEADERS = (
    "Code",
    "Branch",
    "Format",
    "Revenue",
    "Revenue last week",
    "WoW",
    "Transactions",
    "Average basket",
    "Waste EUR",
    "Waste %",
    "Labour hours",
    "Labour cost",
    "Labour % of revenue",
    "Revenue per labour hour",
    "Rank",
    "Rank change",
)
METRIC_COLUMNS = {
    "revenue": "D",
    "transactions": "G",
    "avg_basket": "H",
    "waste_eur": "I",
    "waste_pct": "J",
    "labour_hours": "K",
    "labour_cost": "L",
    "labour_pct": "M",
    "revenue_per_labour_hour": "N",
}
TREND_HEADERS = (
    "Week",
    "Week start",
    "Revenue",
    "Transactions",
    "Waste %",
    "Labour % of revenue",
    "Branches reporting",
)
TREND_COLUMNS = {"revenue": "C", "transactions": "D", "waste_pct": "E", "labour_pct": "F"}
WIDTHS = {
    "A": 9,
    "B": 16,
    "C": 15,
    "D": 13,
    "E": 13,
    "F": 10,
    "G": 13,
    "H": 11,
    "I": 11,
    "J": 9,
    "K": 11,
    "L": 12,
    "M": 11,
    "N": 12,
    "O": 10,
    "P": 9,
}


@dataclass(frozen=True)
class Layout:
    first: int
    last: int
    total: int
    trend_title: int
    trend_header: int
    trend_first: int
    trend_last: int


def layout(cfg: Config) -> Layout:
    last = 3 + len(cfg.branches)
    total = last + 1
    return Layout(4, last, total, total + 2, total + 3, total + 4, total + 3 + cfg.model.trend_weeks)


def _column(ref: str) -> str:
    return ref.rstrip("0123456789")


def total_cells(cfg: Config, result: MetricsResult) -> list[FormulaCell]:
    """The total row: SUM for sum metrics, the metric expressions for ratios, and the revenue change."""
    lay = layout(cfg)
    row = lay.total
    period, previous = result.period, result.previous

    def ref(name: str) -> str:
        return f"{METRIC_COLUMNS[name]}{row}"

    cells = []
    for key, column in METRIC_COLUMNS.items():
        formula = styles.metric_formula(cfg.metrics[key], column, lay.first, lay.last, ref)
        if formula is not None:
            cells.append(FormulaCell(SHEET, f"{column}{row}", formula, result.chain_value(period, key)))
    kind = cfg.metric("revenue").change
    current, before = result.chain_value(period, "revenue"), result.chain_value(previous, "revenue")
    cells.append(FormulaCell(SHEET, f"E{row}", f"=SUM(E{lay.first}:E{lay.last})", before))
    cells.append(
        FormulaCell(
            SHEET, f"F{row}", "=" + change_excel(f"D{row}", f"E{row}", kind), change(current, before, kind)
        )
    )
    return sorted(cells, key=lambda cell: column_index_from_string(_column(cell.ref)))


def _formatting(ws: Any, cfg: Config, lay: Layout) -> None:
    rows = (lay.first, lay.last)
    ws.conditional_formatting.add(
        f"D{rows[0]}:D{rows[1]}", DataBarRule(start_type="min", end_type="max", color=styles.ACCENT)
    )
    for rule in styles.change_rules(cfg.metric("revenue").better):
        ws.conditional_formatting.add(f"F{rows[0]}:F{rows[1]}", rule)
    for threshold in cfg.model.thresholds:
        if threshold.metric in METRIC_COLUMNS:
            column = METRIC_COLUMNS[threshold.metric]
            ws.conditional_formatting.add(
                f"{column}{rows[0]}:{column}{rows[1]}",
                styles.threshold_rule(threshold.op, threshold.value, threshold.style),
            )
        if threshold.metric in TREND_COLUMNS:
            column = TREND_COLUMNS[threshold.metric]
            ws.conditional_formatting.add(
                f"{column}{lay.trend_first}:{column}{lay.trend_last}",
                styles.threshold_rule(threshold.op, threshold.value, threshold.style),
            )


def write(ws: Any, ctx: BuildContext) -> None:
    cfg, result = ctx.cfg, ctx.result
    period, previous = result.period, result.previous
    lay = layout(cfg)
    ranking = {row.code: row for row in ctx.summary.ranking}
    revenue = cfg.metric("revenue")
    styles.title(ws, "A1", f"Branches · week {period}")
    styles.header(ws, 3, 1, HEADERS)
    for row, branch in enumerate(cfg.branches, start=lay.first):
        ws.cell(row=row, column=1, value=branch.code)
        ws.cell(row=row, column=2, value=branch.name)
        ws.cell(row=row, column=3, value=branch.format)
        before = result.branch_value(previous, branch.code, "revenue")
        styles.put(ws, row, 5, before, styles.NUMBER_FORMATS[revenue.format])
        values = result.branch.get((period, branch.code))
        if values is None:
            ws.cell(row=row, column=15, value="no return")
            continue
        for key, column in METRIC_COLUMNS.items():
            number_format = styles.NUMBER_FORMATS[cfg.metric(key).format]
            styles.put(ws, row, column_index_from_string(column), values.get(key), number_format)
        styles.put(
            ws,
            row,
            6,
            change(values.get("revenue"), before, revenue.change),
            styles.CHANGE_FORMATS[revenue.change],
        )
        styles.put(ws, row, 15, ranking[branch.code].rank, "0")
        styles.put(ws, row, 16, ranking[branch.code].rank_change, styles.RANK_CHANGE_FORMAT)

    ws.cell(row=lay.total, column=1, value="Total")
    formats = {
        column: styles.NUMBER_FORMATS[cfg.metric(key).format] for key, column in METRIC_COLUMNS.items()
    }
    formats["E"] = styles.NUMBER_FORMATS[revenue.format]
    formats["F"] = styles.CHANGE_FORMATS[revenue.change]
    for cell in total_cells(cfg, result):
        target = ws[cell.ref]
        target.value = cell.formula
        target.number_format = formats[_column(cell.ref)]
    styles.total_row(ws, lay.total, 1, len(HEADERS))

    ws.cell(
        row=lay.trend_title, column=1, value=f"Weekly totals, last {len(result.window)} weeks"
    ).font = styles.BOLD_FONT
    styles.header(ws, lay.trend_header, 1, TREND_HEADERS)
    for row, week in enumerate(result.window, start=lay.trend_first):
        ws.cell(row=row, column=1, value=str(week))
        styles.put(ws, row, 2, week.monday, styles.DATE_FORMAT)
        for key, column in TREND_COLUMNS.items():
            number_format = styles.NUMBER_FORMATS[cfg.metric(key).format]
            styles.put(
                ws, row, column_index_from_string(column), result.chain_value(week, key), number_format
            )
        styles.put(ws, row, 7, result.reporting.get(week, 0), styles.NUMBER_FORMATS["count"])

    _formatting(ws, cfg, lay)
    styles.widths(ws, WIDTHS)
    ws.row_dimensions[3].height = 30
    ws.row_dimensions[lay.trend_header].height = 30
    ws.freeze_panes = "C4"
    styles.page(
        ws, period, landscape=True, fit_height=1, print_area=f"A1:P{lay.trend_last}", title_rows="3:3"
    )
