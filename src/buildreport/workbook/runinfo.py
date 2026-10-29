"""Sheet Run info: run facts, the input files with their status, and the findings with severity fills."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Any

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font

from buildreport import __version__
from buildreport.isoweek import format_range
from buildreport.metrics import format_value
from buildreport.reproducible import environment
from buildreport.workbook import styles

if TYPE_CHECKING:
    from buildreport.workbook import BuildContext

SHEET = "Run info"
KEYS = (
    "Period",
    "Period dates",
    "Generated at",
    "Timestamp source",
    "Tool",
    "Python",
    "openpyxl",
    "pydantic",
    "PyYAML",
    "Configuration file",
    "Configuration SHA-256",
    "Strict mode",
    "Input folder",
    "Thresholds",
)
INPUT_HEADERS = ("File", "Branch", "Week", "Status", "Submitted at", "Bytes", "SHA-256")
FINDING_HEADERS = ("Severity", "Code", "File", "Message")
WIDTHS = {"A": 28, "B": 22, "C": 24, "D": 14, "E": 17, "F": 9, "G": 68}


@dataclass(frozen=True)
class Layout:
    inputs_title: int
    inputs_header: int
    inputs_first: int
    inputs_last: int
    findings_title: int
    findings_header: int
    findings_first: int
    findings_last: int


def layout(ctx: BuildContext) -> Layout:
    inputs_title = 3 + len(KEYS) + 1
    inputs_last = inputs_title + 1 + len(ctx.resolution.inputs)
    findings_title = inputs_last + 2
    findings_first = findings_title + 2
    return Layout(
        inputs_title=inputs_title,
        inputs_header=inputs_title + 1,
        inputs_first=inputs_title + 2,
        inputs_last=inputs_last,
        findings_title=findings_title,
        findings_header=findings_title + 1,
        findings_first=findings_first,
        findings_last=findings_first + max(len(ctx.resolution.findings), 1) - 1,
    )


def _thresholds(ctx: BuildContext) -> str:
    parts = []
    for threshold in ctx.cfg.model.thresholds:
        spec = ctx.cfg.metric(threshold.metric)
        parts.append(
            f"{spec.label} {threshold.op} {format_value(threshold.value, spec.format)} ({threshold.style})"
        )
    return "; ".join(parts) or "none"


def shown_path(path: Path) -> str:
    """The path relative to the working directory, or only its name when it lies outside it."""
    full = path.resolve()
    try:
        return full.relative_to(Path.cwd().resolve()).as_posix()
    except ValueError:
        return full.name


def key_values(ctx: BuildContext) -> list[tuple[str, str]]:
    env = environment()
    period = ctx.result.period
    values = [
        str(period),
        format_range(period),
        ctx.timestamp.iso,
        ctx.timestamp.source,
        f"build-report {__version__}",
        env["python"],
        env["openpyxl"],
        env["pydantic"],
        env["PyYAML"],
        shown_path(ctx.cfg.path),
        ctx.cfg.sha256,
        "yes" if ctx.strict else "no",
        shown_path(Path(ctx.input_dir)),
        _thresholds(ctx),
    ]
    return list(zip(KEYS, values, strict=True))


def write(ws: Any, ctx: BuildContext) -> None:
    lay = layout(ctx)
    styles.title(ws, "A1", "Run info")
    for row, (key, value) in enumerate(key_values(ctx), start=3):
        ws.cell(row=row, column=1, value=key).font = styles.BOLD_FONT
        ws.cell(row=row, column=2, value=value)

    ws.cell(row=lay.inputs_title, column=1, value="Input files").font = styles.BOLD_FONT
    styles.header(ws, lay.inputs_header, 1, INPUT_HEADERS)
    for row, item in enumerate(ctx.resolution.inputs, start=lay.inputs_first):
        ws.cell(row=row, column=1, value=item.file)
        ws.cell(row=row, column=2, value=item.branch)
        ws.cell(row=row, column=3, value=item.week)
        ws.cell(row=row, column=4, value=item.status)
        ws.cell(row=row, column=5, value=item.submitted_at)
        styles.put(ws, row, 6, item.bytes, "#,##0")
        ws.cell(row=row, column=7, value=item.sha256)

    ws.cell(row=lay.findings_title, column=1, value="Findings").font = styles.BOLD_FONT
    styles.header(ws, lay.findings_header, 1, FINDING_HEADERS)
    findings = ctx.resolution.findings
    if not findings:
        ws.cell(row=lay.findings_first, column=1, value="No findings.")
    for row, item in enumerate(findings, start=lay.findings_first):
        ws.cell(row=row, column=1, value=item.severity)
        ws.cell(row=row, column=2, value=item.code)
        ws.cell(row=row, column=3, value=item.file or "-")
        ws.cell(row=row, column=4, value=item.message)
    if findings:
        first, last = lay.findings_first, lay.findings_last
        target = f"A{first}:D{last}"
        ws.conditional_formatting.add(
            target,
            FormulaRule(
                formula=[f'$A{first}="error"'], fill=styles.bad_fill(), font=Font(color=styles.BAD_TEXT)
            ),
        )
        ws.conditional_formatting.add(
            target,
            FormulaRule(
                formula=[f'$A{first}="warning"'], fill=styles.warn_fill(), font=Font(color=styles.WARN_TEXT)
            ),
        )

    styles.widths(ws, WIDTHS)
    styles.page(ws, ctx.result.period, landscape=False, fit_height=0)
