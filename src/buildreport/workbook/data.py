"""Sheet Data: the long records as the Excel table tblWeekly."""

from __future__ import annotations

from typing import TYPE_CHECKING, Any

from openpyxl.worksheet.table import Table, TableStyleInfo

from buildreport.isoweek import IsoWeek
from buildreport.resolve import Record
from buildreport.workbook import styles

if TYPE_CHECKING:
    from buildreport.workbook import BuildContext

SHEET = "Data"
TABLE = "tblWeekly"
HEADERS = ("week", "week_start", "branch_code", "branch_name", "category", "measure", "value", "source_file")
WIDTHS = {"A": 10, "B": 12, "C": 12, "D": 14, "E": 24, "F": 17, "G": 13, "H": 28}


def write(ws: Any, ctx: BuildContext) -> None:
    names = {branch.code: branch.name for branch in ctx.cfg.branches}
    ws.append(list(HEADERS))
    for record in ctx.records:
        ws.append(
            [
                str(record.week),
                record.week.monday,
                record.branch_code,
                names.get(record.branch_code, ""),
                record.category or None,
                record.measure,
                record.value,
                record.source_file,
            ]
        )
    last = len(ctx.records) + 1
    for row in range(2, last + 1):
        ws.cell(row=row, column=2).number_format = styles.DATE_FORMAT
        ws.cell(row=row, column=7).number_format = styles.VALUE_FORMAT
    table = Table(displayName=TABLE, ref=f"A1:H{last}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    styles.widths(ws, WIDTHS)
    ws.freeze_panes = "A2"
    styles.page(ws, ctx.result.period, landscape=True, fit_height=0, title_rows="1:1", gridlines=True)


def read_records(book: Any) -> list[Record]:
    """The Data sheet back into records; rows without a week or a numeric value are skipped."""
    records = []
    for week, _start, code, _name, category, measure, value, source in book[SHEET].iter_rows(
        min_row=2, max_col=8, values_only=True
    ):
        if week is None or measure is None or isinstance(value, bool) or not isinstance(value, (int, float)):
            continue
        records.append(
            Record(
                IsoWeek.parse(str(week)),
                str(code),
                str(category or ""),
                str(measure),
                float(value),
                str(source or ""),
            )
        )
    return records
